"""
Neck catalog builder — same template as back/chest/biceps but with NECK exercise
science. Neck training MUST stay light and controlled, so the level logic caps at
Intermediate and every `instructions` array (written to be READ ALOUD by the app's
speaker/TTS feature) stresses SLOW, CONTROLLED, LIGHT load through a pain-free range.
Run from this folder: python build_neck_catalog.py
"""
import os, re, json, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Neck_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "neck_catalog.json")
TITLE = "Neck"


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)
    name_part = s.split('_')[0]
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender, 'bodypart': 'Neck'}


def equipment(n):
    l = n.lower()
    if 'harness' in l: return 'Head Harness'
    if 'cable' in l: return 'Cable'
    if 'barbell' in l: return 'Barbell'
    if 'dumbbell' in l: return 'Dumbbell'
    if 'plate' in l or 'weighted' in l: return 'Weight Plate'
    if 'ball' in l: return 'Massage Ball'
    if re.search(r'\bband\b', l) or 'resistance band' in l: return 'Resistance Band'
    if 'wall' in l: return 'Wall'
    return 'Bodyweight'


def movement(n):
    l = n.lower()
    # Stretch / mobility / self-myofascial (foam-rolling). `\broll\b` catches
    # "Roll Neck Rotation" etc. but NOT "Rollout" (one word = no boundary).
    if any(k in l for k in ['stretch', 'mobility', 'foam']) or re.search(r'\broll\b', l):
        return 'Stretch/Mobility'
    if 'lateral' in l or 'side' in l: return 'Lateral Flexion'
    if 'rotation' in l or 'rotate' in l or 'turn' in l: return 'Rotation'
    if 'extension' in l or 'back' in l: return 'Extension'
    if 'flexion' in l or 'front' in l or 'curl' in l or 'chin tuck' in l: return 'Flexion'
    if 'isometric' in l or 'hold' in l: return 'Isometric'
    if 'harness' in l or 'shrug' in l: return 'Extension'
    return 'Flexion'


CATEGORY = {  # movement -> training category (folder name, NO "/" — use "&")
    'Flexion': 'Flexion (Front)',
    'Extension': 'Extension (Back)',
    'Lateral Flexion': 'Lateral & Rotation',
    'Rotation': 'Lateral & Rotation',
    'Isometric': 'Isometric Holds',
    'Stretch/Mobility': 'Stretch & Mobility',
}

# harness / weighted plate -> Intermediate; everything else (stretch, isometric,
# bodyweight) -> Beginner. NOTHING is Advanced — neck training stays light.
WEIGHTED = ['harness', 'weighted', 'plate', 'cable']


def level(n, mv):
    l = n.lower()
    if any(k in l for k in WEIGHTED):
        return 'Intermediate'
    return 'Beginner'


def muscles(mv):
    return {
        'Flexion':         ('Sternocleidomastoid & deep neck flexors', 'Scalenes'),
        'Extension':       ('Splenius capitis & Upper trapezius', 'Levator scapulae, Suboccipitals'),
        'Lateral Flexion': ('Sternocleidomastoid & Scalenes', 'Levator scapulae, Upper trapezius'),
        'Rotation':        ('Sternocleidomastoid & Splenius', 'Scalenes'),
        'Isometric':       ('Cervical flexors & extensors', 'Scalenes, Trapezius'),
        'Stretch/Mobility':('Neck flexors & extensors (stretch)', 'Upper trapezius, Levator scapulae'),
    }.get(mv, ('Sternocleidomastoid & deep neck flexors', 'Scalenes'))


def describe(eq, mv):
    rl = {
        'Flexion': "a controlled neck flexion that strengthens the front of the neck — the sternocleidomastoid and deep neck flexors — by drawing the chin toward the chest.",
        'Extension': "a controlled neck extension that strengthens the back of the neck — the splenius and upper traps — by looking up and back through a comfortable range.",
        'Lateral Flexion': "a controlled side-bend of the neck that works the sternocleidomastoid and scalenes by lowering the ear toward the shoulder.",
        'Rotation': "a controlled neck rotation that works the sternocleidomastoid and splenius by smoothly turning the head to look over the shoulder.",
        'Isometric': "a static neck hold that builds stability in the cervical flexors and extensors by pressing gently and holding without movement.",
        'Stretch/Mobility': "a gentle neck stretch and mobility drill to lengthen the neck flexors and extensors and ease tension in the upper traps.",
    }.get(mv, "a controlled, light movement for the neck.")
    art = 'An ' if rl[0] in 'AEIOUaeiou' else 'A '
    eqs = '' if eq == 'Bodyweight' else f" Performed with {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}."
    return art + rl + eqs


def instructions(eq, mv):
    """Step-by-step how-to, written to be READ ALOUD (speaker/TTS). Returns a list.
    EVERY set stresses SLOW, CONTROLLED, LIGHT load and a pain-free range, and closes
    with a safety cue: never jerk or use heavy load on the neck."""
    SAFE = "Keep every rep slow and controlled, and never jerk the head or use heavy load on the neck."
    STEPS = {
        'Flexion': [
            "Sit or lie down tall with your shoulders relaxed and your head in a neutral position.",
            "Slowly tuck your chin and lower it toward your chest, moving only as far as feels comfortable.",
            "Pause for a moment at the bottom, feeling the front of your neck work.",
            "Return your head smoothly to the start under full control. " + SAFE,
        ],
        'Extension': [
            "Sit or lie down tall with your shoulders relaxed and your head in a neutral position.",
            "Slowly look up and back, taking your head through a comfortable, pain-free range only.",
            "Pause briefly at the top, feeling the back of your neck work.",
            "Lower your head smoothly back to neutral under full control. " + SAFE,
        ],
        'Lateral Flexion': [
            "Sit or stand tall with your shoulders down and relaxed.",
            "Slowly lower one ear toward that same shoulder without lifting the shoulder up to meet it.",
            "Pause for a moment when you feel a gentle stretch and contraction along the side of your neck.",
            "Bring your head smoothly back to center under control, then repeat to the other side. " + SAFE,
        ],
        'Rotation': [
            "Sit or stand tall with your chin level and your shoulders relaxed.",
            "Slowly turn your head to one side as if to look over your shoulder, staying within a comfortable range.",
            "Pause briefly at the end of the turn without forcing it any further.",
            "Return your head smoothly to facing forward, then turn slowly to the other side. " + SAFE,
        ],
        'Isometric': [
            "Sit or stand tall and place a hand against your head in the direction you will press.",
            "Press your head gently into your hand and hold the position without letting your head actually move.",
            "Hold steadily for the prescribed time, breathing slowly and keeping the effort light.",
            "Release the pressure gradually and relax. " + SAFE,
        ],
        'Stretch/Mobility': [
            "Sit or stand tall and ease your head slowly into the stretch until you feel a gentle, comfortable pull.",
            "Hold the position and breathe deeply, letting the neck muscles lengthen with each exhale.",
            "Come out of the stretch slowly and repeat on the other side if needed. " + SAFE,
        ],
    }
    return STEPS.get(mv, [
        "Sit or stand tall with your head in a neutral position and your shoulders relaxed.",
        "Move your head slowly and smoothly through a comfortable, pain-free range.",
        "Return to the start under full control. " + SAFE,
    ])


def folder_for(gender, mv):
    cat = CATEGORY.get(mv, 'Flexion (Front)')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 - already organized.'); return
    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['category'] = CATEGORY.get(r['movement'], 'Flexion (Front)')
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'])
        r['flag'] = ''
        r['folder'] = folder_for(r['gender'], r['movement'])
        recs.append(r)
    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Primary Muscle',
            'Secondary Muscles', 'Description', 'Step-by-step Instructions (TTS)', 'Variants', 'ID', 'Video File(s)', 'Folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r['name'], r['equipment'], r['gender'], r['level'], r['movement'], r['primary'],
                   r['secondary'], r['description'], '\n'.join(f"{i+1}. {s}" for i, s in enumerate(r['instructions'])),
                   r['variants'], r['id'], r['files'], r['folder']])
    ARIAL = 'Arial'
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'), 'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([38, 18, 9, 13, 20, 34, 38, 56, 70, 9, 11, 50, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 6, 7, 8, 9, 12, 13))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises', len(catalog)]); sm.append([])
    for t, k in [('By equipment', 'equipment'), ('By gender', 'gender'), ('By level', 'level'), ('By movement type', 'movement')]:
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    sm.column_dimensions['A'].width = 38; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\nSRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\nprint('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
