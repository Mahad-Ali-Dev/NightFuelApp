"""
Cardio catalog builder — same template as back/chest/biceps but with CARDIO
exercise science. Cardio is NOT a single muscle: exercises are categorized by
MODALITY (run / cycle / row / jump rope / plyometric / HIIT / low-impact /
stretch) and target the CARDIOVASCULAR SYSTEM. Includes a step-by-step
`instructions` array per exercise written to be READ ALOUD by the app's
speaker/TTS feature, with a safety/pacing cue per modality.
Run from this folder: python build_cardio_catalog.py
"""
import os, re, json, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Cardio_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "cardio_catalog.json")
TITLE = "Cardio"


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)
    name_part = s.split('_')[0]
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender, 'bodypart': 'Cardio'}


def equipment(n):
    l = n.lower()
    if 'treadmill' in l: return 'Treadmill'
    if 'elliptical' in l: return 'Elliptical'
    if 'rowing' in l or 'erg' in l or 'rowing machine' in l: return 'Rowing Machine'
    if 'assault' in l or 'air bike' in l or 'hands bike' in l or 'hand bike' in l: return 'Air Bike'
    if 'stationary bike' in l or 'recline' in l or 'outdoor bicycle' in l or 'bicycle' in l or 'bike' in l: return 'Bike'
    if 'stepper' in l or 'lever step' in l or 'lever stepper' in l or 'step ' in l or 'wave machine' in l or 'wave-machine' in l: return 'Stepper'
    if 'battling rope' in l or 'battle rope' in l: return 'Battle Ropes'
    if 'jump rope' in l or 'skip' in l: return 'Jump Rope'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'dumbbell' in l: return 'Dumbbell'
    if re.search(r'\bband\b', l) or 'resistance band' in l: return 'Resistance Band'
    return 'Bodyweight'


def movement(n):
    l = n.lower()
    # Stretch / mobility / self-myofascial (foam-rolling). `\broll\b` catches
    # "Roll Out Hamstring" etc. but NOT "Rollout" (one word = no boundary).
    if any(k in l for k in ['stretch', 'mobility', 'opener', 'foam', 'roller', 'release']) or re.search(r'\broll\b', l):
        return 'Stretch/Mobility'
    if any(k in l for k in ['run', 'treadmill', 'sprint', 'jog']): return 'Run'
    if any(k in l for k in ['cycle', 'bike', 'spin']): return 'Cycle'
    if 'row' in l or 'erg' in l: return 'Row'
    if 'rope' in l or 'skip' in l: return 'Jump Rope'
    if any(k in l for k in ['burpee', 'box jump', 'tuck jump', 'broad jump', 'jump squat', 'plyo']): return 'Plyometric'
    if any(k in l for k in ['mountain climber', 'jumping jack', 'high knee', 'squat thrust', 'shuffle', 'skater']): return 'HIIT'
    if any(k in l for k in ['elliptical', 'step', 'stair', 'march', 'walk']): return 'Low-Impact'
    return 'HIIT'


# movement -> training category (folder, NO "/" — use "&")
CATEGORY = {
    'Run': 'Machine & Steady-State', 'Cycle': 'Machine & Steady-State',
    'Row': 'Machine & Steady-State', 'Low-Impact': 'Machine & Steady-State',
    'Jump Rope': 'Plyometric & Jump', 'Plyometric': 'Plyometric & Jump',
    'HIIT': 'HIIT & Bodyweight', 'Stretch/Mobility': 'Stretch & Mobility',
}


def level(n, mv):
    l = n.lower()
    if mv == 'Stretch/Mobility' or any(k in l for k in ['walk', 'march']) or mv == 'Low-Impact':
        return 'Beginner'
    if any(k in l for k in ['sprint', 'burpee', 'box jump', 'plyo', 'tuck']):
        return 'Advanced'
    return 'Intermediate'


def muscles(mv):
    return {
        'Run':              ('Cardiovascular system (full body)', 'Quadriceps, Hamstrings, Calves, Glutes'),
        'Cycle':            ('Cardiovascular system (full body)', 'Quadriceps, Glutes, Calves'),
        'Row':              ('Cardiovascular system (full body)', 'Back, Legs, Core'),
        'Jump Rope':        ('Cardiovascular system (full body)', 'Calves, Shoulders, Core'),
        'Plyometric':       ('Cardiovascular system (full body)', 'Quadriceps, Glutes, Calves, Core'),
        'HIIT':             ('Cardiovascular system (full body)', 'Full body, Core'),
        'Low-Impact':       ('Cardiovascular system (full body)', 'Legs, Glutes'),
        'Stretch/Mobility': ('Full body (stretch)', 'Hips, Hamstrings, Calves'),
    }.get(mv, ('Cardiovascular system (full body)', 'Full body, Core'))


def describe(eq, mv):
    rl = {
        'Run': "steady-state and interval cardio staple that drives heart rate up and builds aerobic endurance through the legs.",
        'Cycle': "low-joint-impact cardio effort that trains the heart and lungs while building leg endurance.",
        'Row': "full-body cardio effort that combines a leg drive, hip hinge and arm pull to tax the heart and lungs.",
        'Jump Rope': "fast, rhythmic cardio drill that spikes heart rate and trains coordination, calves and footwork.",
        'Plyometric': "explosive, high-intensity conditioning movement that elevates heart rate fast and builds power.",
        'HIIT': "high-intensity interval movement that drives the heart rate up quickly with full-body effort.",
        'Low-Impact': "gentle, joint-friendly cardio effort that keeps the heart rate in an aerobic zone.",
        'Stretch/Mobility': "mobility and stretch drill to warm up or cool down the body around cardio work.",
    }.get(mv, "conditioning movement that trains the cardiovascular system.")
    art = 'An ' if rl[0] in 'AEIOUaeiou' else 'A '
    eqs = '' if eq == 'Bodyweight' else f" Performed on {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}." \
        if eq in ('Treadmill', 'Elliptical', 'Rowing Machine', 'Air Bike', 'Bike', 'Stepper') else \
        f" Performed with {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}."
    return art + rl + eqs


def instructions(eq, mv):
    """Step-by-step how-to, written to be READ ALOUD (speaker/TTS). Returns a list.
    Each modality carries a SAFETY/pacing cue and emphasizes controlled breathing."""
    STEPS = {
        'Run': [
            "Begin with a light warm-up walk or easy jog for two to three minutes to raise your heart rate gradually.",
            "Build to your target pace, maintaining an upright posture and a steady rhythm.",
            "Land midfoot under your hips, drive your arms front to back, and breathe steadily in through your nose and out through your mouth.",
            "Hold your effort for the planned interval, keeping your shoulders relaxed and your stride smooth.",
            "Ease back to a walk to recover. Do not sprint from a cold start, and keep your breathing controlled the whole time.",
        ],
        'Cycle': [
            "Set the seat so your knee is only slightly bent at the bottom of the pedal stroke, then start pedalling easily.",
            "Settle into a smooth, even cadence with a slight forward lean and a relaxed grip.",
            "Increase resistance or speed for your work interval, pushing and pulling through the full pedal circle.",
            "Keep your cadence smooth and your breathing rhythmic — exhale on effort, inhale on recovery.",
            "Spin easy to cool down. Keep your knees tracking forward and never let resistance jerk you out of the saddle.",
        ],
        'Row': [
            "Sit tall, strap in your feet, and grip the handle with your arms straight and shins vertical.",
            "Drive through your legs first, then swing your hips back, and finally pull the handle to your lower ribs.",
            "Reverse the order to return: extend your arms, hinge forward from the hips, then bend your knees to slide forward.",
            "Keep a strong flat back and a steady stroke rate, breathing out on the drive and in on the recovery.",
            "Slow the stroke to cool down. Lead with the legs, not the lower back, and avoid rounding your spine.",
        ],
        'Jump Rope': [
            "Hold the handles at hip height with your elbows close to your sides and the rope behind your heels.",
            "Turn the rope using small wrist turns, not big arm swings.",
            "Take light bounces on the balls of your feet, clearing the rope by just an inch or two.",
            "Find a steady rhythm and keep your breathing relaxed and even.",
            "Slow your turns to stop. Stay tall, keep the bounces low and soft, and land lightly to protect your knees.",
        ],
        'Plyometric': [
            "Start standing tall with your feet hip-width apart and your core braced.",
            "Drop quickly into a quarter squat, then explode upward or outward with full effort.",
            "For a burpee, lower your chest to the floor, jump your feet in, then explode up; reset each rep with control.",
            "Absorb every landing softly through bent knees and hips, breathing out on the effort.",
            "Move at a hard but sustainable intensity. Land softly, keep your knees tracking over your toes, and stop if your form breaks down.",
        ],
        'HIIT': [
            "Set up in the start position with your core tight and your weight balanced.",
            "Work explosively for the interval — drive the knees fast for mountain climbers, or jump wide and back for jumping jacks.",
            "Keep your hips low and level on floor moves and stay light on your feet on standing moves.",
            "Push hard through the work period while keeping your breathing as controlled as you can.",
            "Rest fully between bouts. Keep your movements crisp, and back off the intensity before your technique falls apart.",
        ],
        'Low-Impact': [
            "Start at an easy effort and let your heart rate climb gradually over the first minute.",
            "Settle into a smooth, controlled rhythm with a tall posture and a steady stride or step.",
            "Hold an aerobic pace where you can still breathe in a steady, controlled rhythm.",
            "Keep your movements even and your core engaged throughout the effort.",
            "Wind down slowly to finish. Keep the impact gentle and the intensity moderate to stay in the aerobic zone.",
        ],
        'Stretch/Mobility': [
            "Move into the stretch slowly until you feel a gentle pull through the target muscle.",
            "Hold the position and breathe deeply, letting the muscle lengthen with each exhale.",
            "Keep the movement smooth and controlled — never bounce or force the range.",
            "Ease out of the stretch under control and repeat on the other side if needed.",
        ],
    }
    return STEPS.get(mv, [
        "Set up in a balanced, athletic stance with your core braced.",
        "Work at a steady, controlled intensity for the planned interval.",
        "Keep your breathing rhythmic and your movements smooth, then recover fully.",
    ])


def folder_for(gender, mv):
    cat = CATEGORY.get(mv, 'HIIT & Bodyweight')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 - already organized.'); return
    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['category'] = CATEGORY.get(r['movement'], 'HIIT & Bodyweight')
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'])
        r['flag'] = ''
        r['folder'] = folder_for(r['gender'], r['movement'])
        recs.append(r)
    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Primary Muscle',
            'Secondary Muscles', 'Description', 'Step-by-step Instructions (TTS)', 'Variants', 'ID', 'Video File(s)', 'Folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r['name'], r['equipment'], r['gender'], r['level'], r['movement'], r['primary'],
                   r['secondary'], r['description'], '\n'.join(f"{i+1}. {s}" for i, s in enumerate(r['instructions'])),
                   r['variants'], r['id'], r['files'], r['folder']])
    ARIAL = 'Arial'
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'), 'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([38, 18, 9, 13, 20, 34, 38, 56, 70, 9, 11, 50, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 6, 7, 8, 9, 12, 13))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises', len(catalog)]); sm.append([])
    for t, k in [('By equipment', 'equipment'), ('By gender', 'gender'), ('By level', 'level'), ('By movement type', 'movement')]:
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    sm.column_dimensions['A'].width = 38; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\nSRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\nprint('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
