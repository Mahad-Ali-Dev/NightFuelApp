import os, re, json, sys, shutil
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"C:\Users\Sajjad Ali\Downloads\chest"
OUT_XLSX = os.path.join(SRC, "Chest_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "chest_catalog.json")

BODYPARTS = {'chest':'Chest','ches':'Chest','ch':'Chest','c':'Chest','back':'Back','waist':'Waist',
 'shoulders':'Shoulders','hips':'Hips','upper-arms':'Upper Arms','plyometrics':'Plyometrics',
 'stretching':'Stretching','neck':'Neck'}

def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)
    parts = s.split('_')
    name_part = parts[0]
    bodypart = ''
    for t in parts[1:]:
        t2 = re.sub(r'-?FIX2?', '', t, flags=re.I).strip(' -')
        if t2:
            bodypart = t2; break
    low = name_part.lower().strip()
    if 'female' in low or re.search(r'\(fe', low) or low.endswith('(f'):
        gender = 'Female'
    elif '(male' in low or low.endswith('-m') or '-m_' in s.lower():
        gender = 'Male'
    else:
        gender = 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'\(fe?$', '', name.strip())
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'(?<=\w)-m$', '', name.strip())
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender,
            'bodypart': BODYPARTS.get(bodypart.lower(), bodypart or 'Chest')}

def equipment(n):
    l = n.lower()
    if 'landmine' in l: return 'Landmine'
    if 'smith' in l: return 'Smith Machine'
    if 'barbell' in l: return 'Barbell'
    if 'dumbbell' in l: return 'Dumbbell'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'cable' in l: return 'Cable'
    if 'pec deck' in l or 'lever' in l or 'plate-loaded' in l or 'plate loaded' in l or 'machine' in l: return 'Machine (Lever)'
    if 'resistance band' in l or re.search(r'\bband\b', l): return 'Resistance Band'
    if 'suspender' in l or 'suspension' in l or 'trx' in l: return 'Suspension (TRX)'
    if 'ring' in l: return 'Rings'
    if 'medicine ball' in l: return 'Medicine Ball'
    if 'battling rope' in l or 'battle rope' in l: return 'Battle Ropes'
    if 'bottle' in l: return 'Improvised (bottle)'
    if 'plate' in l and ('svend' in l or 'weighted' in l): return 'Weight Plate'
    if 'foam' in l or l.startswith('roll') or 'roller' in l: return 'Foam Roller'
    return 'Bodyweight'

def region(n):
    l = n.lower()
    if 'push' in l and 'up' in l:
        if 'decline' in l: return 'Upper'
        if 'incline' in l: return 'Lower'
    if 'incline' in l or 'high' in l or 'upper' in l: return 'Upper'
    if 'decline' in l or 'low' in l: return 'Lower'
    return 'Mid'

def movement(n):
    l = n.lower()
    if any(k in l for k in ['stretch','opener','mobility','foam','roller','clam','sky look','doorway','release']) or l.startswith('roll'):
        return 'Stretch/Mobility'
    if 'pullover' in l: return 'Pullover'
    if 'dip' in l: return 'Dip'
    if 'svend' in l: return 'Fly / Squeeze'
    if 'squeeze' in l and 'press' not in l and 'bench' not in l: return 'Fly / Squeeze'
    if any(k in l for k in ['fly','flye','crossover','pec deck','adduction']): return 'Fly / Squeeze'
    if any(k in l for k in ['throw','pass','slam','devils press','jab','punch']): return 'Power / Throw'
    if any(k in l for k in ['push-up','push up','pushup','pike','cobra','planche','archer','plyo','clap','shoulder tap','plank','bird dog','crawl']): return 'Push-up'
    if 'press' in l: return 'Press'
    if any(k in l for k in ['circling','flinging','arm circle']): return 'Dynamic warm-up'
    return 'Press'

ADV = ['one-arm','one arm','single-arm','single arm','single leg','one-leg','one leg','one side','archer','planche','pseudo','ring','plyo','clap','explosive','superman','deep push','devils press','slam','typewriter','side-to-side','side to side','rotational','rotating','knuckle','fist','finger','suspended','bear crawl','renegade','staggered','hindu']
BEG = ['wall','on knees','kneeling','knees','assisted','stretch','opener','mobility','foam','roller','isometric','hold','warm-up','warm up','prayer','clam','release','chair','pec deck','machine','lever','band','incline push','circling','flinging','child']

def level(n, mv):
    l = n.lower()
    if any(k in l for k in ADV): return 'Advanced'
    if mv in ('Stretch/Mobility','Dynamic warm-up'): return 'Beginner'
    if any(k in l for k in BEG): return 'Beginner'
    if mv == 'Power / Throw': return 'Advanced'
    return 'Intermediate'

def muscles(mv, reg):
    r = reg.lower()
    if mv == 'Stretch/Mobility': return ('Pectoralis major & minor (stretch)', 'Anterior deltoid')
    if mv == 'Pullover': return ('Pectoralis major', 'Latissimus dorsi, Serratus anterior, Triceps (long head)')
    if mv == 'Fly / Squeeze': return (f'Pectoralis major ({r})', 'Anterior deltoid')
    if mv == 'Dip': return ('Pectoralis major (lower)', 'Triceps brachii, Anterior deltoid')
    if mv == 'Power / Throw': return ('Pectoralis major', 'Anterior deltoid, Triceps brachii, Core')
    if mv == 'Push-up': return (f'Pectoralis major ({r})', 'Anterior deltoid, Triceps brachii, Core')
    if mv == 'Dynamic warm-up': return ('Pectoralis major', 'Anterior deltoid')
    return (f'Pectoralis major ({r})', 'Anterior deltoid, Triceps brachii')

def describe(eq, mv, reg):
    rl = {'Upper':'upper (clavicular) chest','Mid':'mid chest','Lower':'lower chest'}[reg]
    base = {
      'Press': f"press the load away from the chest and lower it under control, emphasizing the {rl}; front delts and triceps assist.",
      'Fly / Squeeze': f"open the arms wide with soft elbows, then squeeze them together across the chest to work the {rl} through adduction.",
      'Dip': "lower the torso between the supports until the chest stretches, then press back up; leaning forward loads the lower chest.",
      'Pullover': "from an overhead stretch, pull the load back over the chest, working the pecs with the lats and serratus.",
      'Push-up': f"hold a rigid plank and lower the chest toward the floor, then press up; trains the {rl}, shoulders, triceps and bracing core.",
      'Power / Throw': "move explosively (throw/clap/drive) to develop chest power — a higher-skill, higher-intensity drill.",
      'Stretch/Mobility': "open and lengthen the chest, holding the position to release the pecs and front shoulders.",
      'Dynamic warm-up': "controlled dynamic movement to warm up and mobilize the chest and shoulders.",
    }[mv]
    if eq == 'Bodyweight':
        return base[0].upper() + base[1:]
    art = 'an ' if eq[0] in 'AEIOU' else 'a '
    return f"Using {art}{eq.lower()}, " + base

def folder_for(gender, mv, reg, flagged):
    if flagged: cat = 'Review (non-chest tag)'
    elif mv == 'Stretch/Mobility': cat = 'Stretch & Mobility'
    elif mv in ('Power / Throw', 'Dynamic warm-up'): cat = 'Other'
    else: cat = {'Upper':'Incline (Upper Chest)','Mid':'Flat (Mid Chest)','Lower':'Decline (Lower Chest)'}[reg]
    return gender + '/' + cat

files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
if not files:
    print('No top-level .mp4 in', SRC, '- already organized. Nothing to do.'); sys.exit(0)

recs = []
for f in files:
    r = parse(f)
    r['equipment'] = equipment(r['name'])
    r['movement'] = movement(r['name'])
    r['region'] = region(r['name'])
    r['level'] = level(r['name'], r['movement'])
    r['primary'], r['secondary'] = muscles(r['movement'], r['region'])
    r['description'] = describe(r['equipment'], r['movement'], r['region'])
    r['flag'] = '' if r['bodypart'] == 'Chest' else 'Review: non-chest tag'
    r['folder'] = folder_for(r['gender'], r['movement'], r['region'], bool(r['flag']))
    recs.append(r)

groups = defaultdict(list)
for r in recs:
    groups[(r['name'].lower(), r['gender'])].append(r)
catalog = []
for grp in groups.values():
    b = dict(grp[0])
    b['variants'] = len(grp)
    b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
    catalog.append(b)
catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))

json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

wb = Workbook(); ws = wb.active; ws.title = 'Chest Catalog'
cols = ['Exercise Name','Equipment','Gender','Level','Movement Type','Primary Muscle',
        'Secondary Muscles','Description','Body-Part Tag','Flag','Variants','ID','Video File(s)','Folder']
ws.append(cols)
keys = ['name','equipment','gender','level','movement','primary','secondary','description','bodypart','flag','variants','id','files','folder']
for r in catalog:
    ws.append([r.get(k, '') for k in keys])

ARIAL = 'Arial'
hdr_fill = PatternFill('solid', fgColor='1F3864')
hdr_font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
thin = Side(style='thin', color='D9D9D9')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'),
            'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
            'Advanced': PatternFill('solid', fgColor='FFC7CE')}
flag_fill = PatternFill('solid', fgColor='FCE4D6')
widths = [40,16,9,13,16,28,34,62,14,20,9,11,52,26]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in range(1, len(cols)+1):
    cell = ws.cell(1, c); cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
ws.row_dimensions[1].height = 30
for row in range(2, ws.max_row+1):
    for c in range(1, len(cols)+1):
        cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
        cell.alignment = Alignment(vertical='top', wrap_text=(c in (1,6,7,8,13,14))); cell.border = border
    lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill()); lv.alignment = Alignment(horizontal='center', vertical='top')
    fl = ws.cell(row, 10)
    if fl.value: fl.fill = flag_fill
    ws.cell(row, 3).alignment = Alignment(horizontal='center', vertical='top')
    ws.cell(row, 11).alignment = Alignment(horizontal='center', vertical='top')
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

sm = wb.create_sheet('Summary')
sm.append(['Chest Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises (deduped)', len(catalog)])
sm.append([])
from collections import Counter
def block(title, key):
    sm.append([title]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
    for val, cnt in sorted(Counter(r[key] for r in catalog).items(), key=lambda x:-x[1]):
        sm.append([val, cnt])
    sm.append([])
block('By equipment', 'equipment'); block('By gender', 'gender'); block('By level', 'level'); block('By movement type', 'movement')
flagged = sum(1 for r in catalog if r['flag'])
sm.append(['Flagged (non-chest tag) for review', flagged]); sm.cell(sm.max_row,1).font = Font(name=ARIAL, bold=True, color='C00000')
sm.column_dimensions['A'].width = 34; sm.column_dimensions['B'].width = 12
for row in sm.iter_rows():
    for cell in row:
        if cell.font.name != ARIAL: cell.font = Font(name=ARIAL, size=10)
wb.save(OUT_XLSX)
print('FILES', len(files), 'UNIQUE', len(catalog), 'FLAGGED', flagged)
print('SAVED', OUT_XLSX)

manifest = []
fc = {}
for f in files:
    r = parse(f)
    rel = folder_for(r['gender'], movement(r['name']), region(r['name']), r['bodypart'] != 'Chest')
    dest_dir = os.path.join(SRC, *rel.split('/'))
    os.makedirs(dest_dir, exist_ok=True)
    src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
    if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
        shutil.move(src, dst)
    manifest.append({'file': f, 'folder': rel})
    fc[rel] = fc.get(rel, 0) + 1
json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
undo = (
"import os, json, shutil\n"
"SRC = os.path.dirname(os.path.abspath(__file__))\n"
"m = json.load(open(os.path.join(SRC, '_organize_manifest.json'), encoding='utf-8'))\n"
"n = 0\n"
"for e in m:\n"
"    cur = os.path.join(SRC, *e['folder'].split('/'), e['file'])\n"
"    if os.path.exists(cur):\n"
"        shutil.move(cur, os.path.join(SRC, e['file'])); n += 1\n"
"for root, dirs, fs in os.walk(SRC, topdown=False):\n"
"    if root != SRC and not os.listdir(root):\n"
"        os.rmdir(root)\n"
"print('Reverted', n, 'files to flat layout')\n"
)
open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(undo)
print('--- MOVED BY FOLDER ---')
for k in sorted(fc): print(f'{fc[k]:4d}  {k}')
print('TOTAL MOVED', sum(fc.values()))
