"""
Biceps catalog builder — mirrors the chest template (parse -> derive metadata ->
dedupe -> JSON + styled XLSX -> organize into Gender/Category folders -> undo),
but with BICEPS-specific exercise science (head emphasis, primary/secondary
muscles, descriptions). Run from this folder: python build_biceps_catalog.py
"""
import os, re, json, sys, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Biceps_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "biceps_catalog.json")
TITLE = "Biceps"

BODYPARTS = {'upper-arms': 'Upper Arms', 'upper-arm': 'Upper Arms', 'forearms': 'Forearms',
             'forearm': 'Forearms', 'arms': 'Upper Arms'}


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)            # trailing (1)(2) dupes
    parts = s.split('_')
    name_part = parts[0]
    bodypart = ''
    for t in parts[1:]:
        t2 = re.sub(r'-?FIX2?', '', t, flags=re.I).strip(' -')
        t2 = re.sub(r'without-?weight', '', t2, flags=re.I).strip(' -')
        if t2:
            bodypart = t2; break
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender,
            'bodypart': BODYPARTS.get(bodypart.lower(), bodypart or 'Upper Arms')}


def equipment(n):
    l = n.lower()
    if 'ez' in l.split() or 'ez-' in l or l.startswith('ez '): return 'EZ-Bar'
    if 'smith' in l: return 'Smith Machine'
    if 'barbell' in l: return 'Barbell'
    if 'dumbbell' in l: return 'Dumbbell'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'cable' in l: return 'Cable'
    if 'machine' in l or 'lever' in l or 'plate-loaded' in l: return 'Machine (Lever)'
    if re.search(r'\bband\b', l) or 'resistance band' in l: return 'Resistance Band'
    if 'suspension' in l or 'trx' in l: return 'Suspension (TRX)'
    return 'Bodyweight'


def movement(n):
    l = n.lower()
    if any(k in l for k in ['stretch', 'opener', 'mobility', 'foam', 'release', 'doorway']):
        return 'Stretch/Mobility'
    if 'zottman' in l: return 'Zottman Curl'
    if 'reverse' in l: return 'Reverse Curl'
    if 'hammer' in l: return 'Hammer Curl'
    if 'preacher' in l: return 'Preacher Curl'
    if 'spider' in l or 'prone incline' in l: return 'Spider Curl'
    if 'concentration' in l: return 'Concentration Curl'
    if 'drag' in l: return 'Drag Curl'
    if 'incline' in l: return 'Incline Curl'
    if 'curl' in l: return 'Curl'
    return 'Curl'


# head-emphasis category from the movement
def region(mv):
    return {
        'Incline Curl': 'Long Head',
        'Drag Curl': 'Long Head',
        'Preacher Curl': 'Short Head',
        'Spider Curl': 'Short Head',
        'Concentration Curl': 'Short Head',
        'Hammer Curl': 'Brachialis',
        'Reverse Curl': 'Brachialis',
        'Zottman Curl': 'Brachialis',
        'Curl': 'Overall',
        'Stretch/Mobility': 'Stretch',
    }.get(mv, 'Overall')


ADV = ['one-arm', 'one arm', 'single-arm', 'single arm', 'zottman', 'drag', 'concentration',
       'spider', 'prone incline', 'suspension', 'trx', 'standing one']
BEG = ['machine', 'lever', 'cable', 'band', 'assisted', 'seated preacher', 'preacher',
       'stretch', 'mobility', 'wall']


def level(n, mv):
    l = n.lower()
    if mv == 'Stretch/Mobility': return 'Beginner'
    if any(k in l for k in ADV): return 'Advanced'
    if any(k in l for k in BEG): return 'Beginner'
    return 'Intermediate'


def muscles(mv):
    return {
        'Curl':               ('Biceps brachii', 'Brachialis, Brachioradialis'),
        'Incline Curl':       ('Biceps brachii (long head)', 'Brachialis'),
        'Drag Curl':          ('Biceps brachii (long head)', 'Brachialis, Posterior deltoid'),
        'Preacher Curl':      ('Biceps brachii (short head)', 'Brachialis'),
        'Spider Curl':        ('Biceps brachii (short head)', 'Brachialis'),
        'Concentration Curl': ('Biceps brachii (short head, peak)', 'Brachialis'),
        'Hammer Curl':        ('Brachialis & Brachioradialis', 'Biceps brachii'),
        'Reverse Curl':       ('Brachioradialis & wrist extensors', 'Biceps brachii, Brachialis'),
        'Zottman Curl':       ('Biceps brachii & Brachioradialis', 'Brachialis, Wrist extensors'),
        'Stretch/Mobility':   ('Biceps brachii (stretch)', 'Brachialis, Anterior forearm'),
    }.get(mv, ('Biceps brachii', 'Brachialis, Brachioradialis'))


def describe(eq, mv):
    base = {
        'Curl': "curl the load from a full stretch to full contraction with palms up, keeping the elbows pinned to the sides and the upper arms still so the biceps do the work.",
        'Incline Curl': "lie back on an incline bench so the arms hang behind the torso; curl from that lengthened position to bias the biceps long (outer) head.",
        'Drag Curl': "drag the load up close to the body, pulling the elbows back rather than forward, to keep tension on the long head through the top.",
        'Preacher Curl': "rest the upper arms on the pad and curl from a deep stretch; the preacher angle kills momentum and loads the short (inner) head.",
        'Spider Curl': "lie chest-down on an incline with the arms hanging straight; curl with no body english for constant short-head tension.",
        'Concentration Curl': "brace the elbow against the inner thigh and curl one arm at a time, squeezing hard at the top for a strong peak contraction.",
        'Hammer Curl': "curl with a neutral (thumbs-up) grip, driving the brachialis and brachioradialis along with the biceps for thicker, stronger arms.",
        'Reverse Curl': "curl with an overhand (palms-down) grip to shift the work onto the brachioradialis and wrist extensors of the forearm.",
        'Zottman Curl': "curl up palms-up, rotate to palms-down at the top, then lower under control — training the biceps on the way up and the forearm on the way down.",
        'Stretch/Mobility': "lengthen and open the biceps and front of the arm, holding the position to release the elbow flexors.",
    }.get(mv, "curl the load with control, keeping the upper arms still to isolate the elbow flexors.")
    if eq == 'Bodyweight':
        return base[0].upper() + base[1:]
    art = 'an ' if eq[0] in 'AEIOU' else 'a '
    return f"Using {art}{eq.lower()}, " + base


def instructions(eq, mv):
    """Step-by-step how-to written to be READ ALOUD by the app's speaker/TTS.
    Returns a list of spoken steps (same shape/style as the back catalog)."""
    tool = 'the bar' if eq in ('Barbell', 'EZ-Bar', 'Smith Machine') else \
           'the handle' if eq == 'Cable' else 'the band' if eq == 'Resistance Band' else \
           'the dumbbells' if eq == 'Dumbbell' else 'the weight'
    STEPS = {
        'Curl': [
            f"Stand tall holding {tool} with an underhand grip, arms straight and elbows pinned to your sides.",
            f"Curl {tool} up toward your shoulders by bending only at the elbows — keep your upper arms still.",
            "Squeeze your biceps hard at the top.",
            f"Lower {tool} slowly all the way back to a full stretch. Don't swing or lean back to cheat the weight up.",
        ],
        'Incline Curl': [
            "Sit back on an incline bench so your arms hang straight down behind your torso.",
            "Curl the dumbbells up while keeping your upper arms back, feeling the stretch across the biceps.",
            "Squeeze at the top, then lower slowly to a full stretch. The behind-the-body angle biases the long (outer) head.",
        ],
        'Drag Curl': [
            f"Stand holding {tool} against the front of your thighs.",
            f"Drag {tool} straight up your body by pulling your elbows back behind you, keeping it close to your torso.",
            "Squeeze at the top, then lower under control. Pulling the elbows back keeps tension on the long head.",
        ],
        'Preacher Curl': [
            f"Rest the backs of your upper arms flat on the preacher pad and hold {tool} with an underhand grip.",
            f"Curl {tool} up from a full stretch, keeping your upper arms pressed into the pad.",
            "Squeeze at the top, then lower slowly until your arms are almost straight. Never bounce out of the bottom stretch.",
        ],
        'Spider Curl': [
            "Lie chest-down on an incline bench with your arms hanging straight toward the floor.",
            "Curl the dumbbells up with no body movement, keeping the upper arms vertical.",
            "Squeeze hard at the top, then lower slowly for constant tension on the short head.",
        ],
        'Concentration Curl': [
            "Sit down and brace the back of your working upper arm against the inside of your thigh.",
            "Curl the weight up toward your shoulder one arm at a time, keeping the upper arm still.",
            "Squeeze hard at the top for a strong peak contraction, then lower slowly under control.",
        ],
        'Hammer Curl': [
            "Stand tall holding the dumbbells with a neutral, thumbs-up grip and your elbows at your sides.",
            "Curl the dumbbells up keeping your palms facing each other the whole time.",
            "Squeeze at the top, then lower slowly. The neutral grip drives the brachialis and forearm along with the biceps.",
        ],
        'Reverse Curl': [
            f"Stand holding {tool} with an overhand, palms-down grip and your elbows pinned to your sides.",
            f"Curl {tool} up by bending the elbows, keeping your wrists firm and straight.",
            "Squeeze at the top, then lower slowly. The palms-down grip shifts the work onto the brachioradialis and forearm.",
        ],
        'Zottman Curl': [
            "Stand holding the dumbbells with an underhand grip and curl them up to the top.",
            "At the top, rotate your wrists so your palms face down.",
            "Lower the dumbbells slowly in that palms-down position, then rotate back to palms-up at the bottom — biceps on the way up, forearm on the way down.",
        ],
        'Stretch/Mobility': [
            "Ease into the stretch slowly until you feel a gentle pull through your biceps and the front of your arm.",
            "Hold the position and breathe, letting the elbow flexors lengthen with each exhale.",
            "Come out of the stretch under control and repeat on the other side if needed.",
        ],
    }
    return STEPS.get(mv, [
        f"Hold {tool} with control and your elbows pinned to your sides.",
        f"Curl {tool} up by bending only at the elbows and squeeze your biceps at the top.",
        f"Lower {tool} slowly to a full stretch and repeat.",
    ])


def folder_for(gender, mv, reg):
    if mv == 'Stretch/Mobility': cat = 'Stretch & Mobility'
    else:
        cat = {'Long Head': 'Long Head (Incline & Drag)',
               'Short Head': 'Short Head (Preacher & Concentration)',
               'Brachialis': 'Brachialis & Forearm (Hammer & Reverse)',
               'Overall': 'Standard Curl'}.get(reg, 'Other')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 in', SRC, '- already organized.'); return

    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['region'] = region(r['movement'])
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'])
        r['flag'] = '' if r['bodypart'] in ('Upper Arms', 'Forearms') else 'Review: non-arm tag'
        r['folder'] = folder_for(r['gender'], r['movement'], r['region'])
        recs.append(r)

    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    # ---- styled xlsx ----
    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Primary Muscle',
            'Secondary Muscles', 'Description', 'Body-Part Tag', 'Flag', 'Variants', 'ID', 'Video File(s)', 'Folder',
            'Step-by-step Instructions (TTS)']
    keys = ['name', 'equipment', 'gender', 'level', 'movement', 'primary', 'secondary',
            'description', 'bodypart', 'flag', 'variants', 'id', 'files', 'folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r.get(k, '') for k in keys] + ['\n'.join(f"{i+1}. {s}" for i, s in enumerate(r.get('instructions', [])))])
    ARIAL = 'Arial'
    hdr_fill = PatternFill('solid', fgColor='1F3864')
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'),
                'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([40, 16, 9, 13, 18, 30, 34, 64, 14, 20, 9, 11, 52, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = hdr_fill
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 6, 7, 8, 13, 14))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises (deduped)', len(catalog)]); sm.append([])
    def block(t, k):
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    block('By equipment', 'equipment'); block('By gender', 'gender'); block('By level', 'level'); block('By movement type', 'movement')
    sm.column_dimensions['A'].width = 36; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    # ---- organize files into Gender/Category ----
    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']), region(movement(r['name'])))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\n"
        "SRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\n"
        "print('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
