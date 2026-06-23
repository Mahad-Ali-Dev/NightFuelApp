"""
Hips & Glutes catalog builder — same template as back/chest/biceps but with
HIPS/GLUTES exercise science (gluteus maximus / medius & minimus / adductors /
hamstrings), and a step-by-step `instructions` array per exercise written to be
READ ALOUD by the app's speaker/TTS feature. Run from this folder:
python build_hips_catalog.py
"""
import os, re, json, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Hips_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "hips_catalog.json")
TITLE = "Hips & Glutes"


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)
    name_part = s.split('_')[0]
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender, 'bodypart': 'Hips'}


def equipment(n):
    l = n.lower()
    if 'smith' in l: return 'Smith Machine'
    if 'trap bar' in l or 'trap-bar' in l or 'landmine' in l: return 'T-Bar / Landmine'
    if 'barbell' in l: return 'Barbell'
    if 'dumbbell' in l or 'dumbell' in l: return 'Dumbbell'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'cable' in l: return 'Cable'
    if 'sled' in l or 'leg press' in l or 'machine' in l or 'lever' in l or 'plate-loaded' in l or 'plate loaded' in l: return 'Machine (Lever)'
    if re.search(r'\bband\b', l) or 'resistance band' in l: return 'Resistance Band'
    if 'suspension' in l or 'trx' in l or 'ring' in l: return 'Suspension (TRX)'
    return 'Bodyweight'


def movement(n):
    l = n.lower()
    # Stretch / mobility / self-myofascial (foam-rolling). `\broll\b` catches
    # "Roll Hip Stretch" etc. but NOT "Rollout" (one word = no boundary = a core move).
    if any(k in l for k in ['stretch', 'mobility', 'foam']) or re.search(r'\broll\b', l):
        return 'Stretch/Mobility'
    if 'thrust' in l: return 'Hip Thrust'
    if 'bridge' in l: return 'Glute Bridge'
    if 'clam' in l or 'fire hydrant' in l: return 'Activation'
    if 'abduction' in l or 'abductor' in l: return 'Abduction'
    if 'adduction' in l or 'adductor' in l or 'inner thigh' in l: return 'Adduction'
    if any(k in l for k in ['kickback', 'donkey kick', 'pull-through', 'pull through',
                            'hip extension', 'extension']):
        return 'Hip Extension'
    if 'lunge' in l: return 'Lunge'
    if 'step up' in l or 'step-up' in l: return 'Step-up'
    if 'squat' in l: return 'Squat'
    return 'Hip Extension'


# movement -> training category (folder). NO "/" in any folder name — use "&".
CATEGORY = {
    'Hip Thrust': 'Glute Bridges & Thrusts',
    'Glute Bridge': 'Glute Bridges & Thrusts',
    'Abduction': 'Abduction (Outer Glutes)',
    'Adduction': 'Adduction (Inner Thigh)',
    'Hip Extension': 'Hip Extension & Kickbacks',
    'Activation': 'Activation & Mobility',
    'Lunge': 'Lunges, Squats & Step-ups',
    'Step-up': 'Lunges, Squats & Step-ups',
    'Squat': 'Lunges, Squats & Step-ups',
    'Stretch/Mobility': 'Stretch & Mobility',
}

ADV = ['single-leg', 'single leg', 'one-leg', 'one leg', 'one-legged', 'one legged',
       'single-arm', 'single arm', 'bulgarian', 'deficit', 'curtsy', 'snatch', 'pause',
       'paused', 'death march', 'jefferson']


def level(n, mv):
    l = n.lower()
    if mv == 'Stretch/Mobility': return 'Beginner'
    # Single-leg / barbell thrust / bulgarian / deficit / curtsy -> Advanced.
    if 'barbell' in l and 'thrust' in l: return 'Advanced'
    if any(k in l for k in ADV): return 'Advanced'
    # machine / band / bridge / clam / activation -> Beginner.
    if mv in ('Glute Bridge', 'Activation'):
        return 'Beginner'
    if any(k in l for k in ['machine', 'lever', 'sled', 'leg press', 'plate-loaded',
                            'plate loaded']) or re.search(r'\bband\b', l) or 'resistance band' in l \
            or 'bridge' in l or 'clam' in l:
        return 'Beginner'
    return 'Intermediate'


def muscles(mv):
    return {
        'Hip Thrust':       ('Gluteus maximus', 'Hamstrings, Quadriceps, Core'),
        'Glute Bridge':     ('Gluteus maximus', 'Hamstrings, Erector spinae'),
        'Abduction':        ('Gluteus medius & minimus', 'Tensor fasciae latae'),
        'Adduction':        ('Hip adductors (adductor magnus, longus, brevis)', 'Gracilis, Pectineus'),
        'Hip Extension':    ('Gluteus maximus', 'Hamstrings, Erector spinae'),
        'Activation':       ('Gluteus medius', 'Gluteus maximus, Deep hip rotators'),
        'Lunge':            ('Gluteus maximus & Quadriceps', 'Hamstrings, Adductors, Core'),
        'Step-up':          ('Gluteus maximus & Quadriceps', 'Hamstrings, Calves'),
        'Squat':            ('Gluteus maximus & Quadriceps', 'Hamstrings, Adductors, Core'),
        'Stretch/Mobility': ('Glutes & Hip flexors (stretch)', 'Piriformis, Adductors'),
    }.get(mv, ('Gluteus maximus', 'Hamstrings, Erector spinae'))


def describe(eq, mv):
    rl = {
        'Hip Thrust': "a glute-dominant hip extension that loads the gluteus maximus hard through a full range — one of the best glute builders.",
        'Glute Bridge': "a floor-based hip extension that activates and strengthens the glutes while teaching a strong, neutral-spine lockout.",
        'Abduction': "a hip-abduction movement that targets the outer glutes — the gluteus medius and minimus — for hip strength and stability.",
        'Adduction': "a hip-adduction movement that works the inner-thigh adductors against resistance.",
        'Hip Extension': "a hip-extension movement that drives the leg back using the glutes and hamstrings while the spine stays neutral.",
        'Activation': "a glute-activation drill that fires up the gluteus medius and deep hip rotators with the hips stacked and braced.",
        'Lunge': "a single-leg pattern that builds the glutes and quads through a long range, with the adductors and core stabilizing.",
        'Step-up': "a single-leg step pattern that drives the glutes and quads, training balance and hip drive.",
        'Squat': "a knee-and-hip-bend pattern that builds the glutes and quads with the whole lower body and core working.",
        'Stretch/Mobility': "a hip stretch/mobility drill to lengthen the glutes, hip flexors and adductors and free up the hips.",
    }.get(mv, "a hip-and-glute movement.")
    art = 'An ' if rl[0] in 'AEIOUaeiou' else 'A '
    eqs = '' if eq == 'Bodyweight' else f" Performed with {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}."
    return art + rl + eqs


def instructions(eq, mv):
    """Step-by-step how-to, written to be READ ALOUD (speaker/TTS). Returns a list."""
    tool = 'the bar' if 'Bar' in eq or eq in ('Barbell', 'Smith Machine', 'T-Bar / Landmine') else \
           'the handle' if eq == 'Cable' else 'the weight' if eq in ('Dumbbell', 'Kettlebell') else 'the resistance'
    STEPS = {
        'Hip Thrust': [
            "Sit on the floor with your upper back resting against a bench and your feet flat, about hip-width apart.",
            f"Position {tool} across your hips and roll it snug against your hip crease.",
            "Tuck your chin, brace your core, and drive through your heels to lift your hips until your torso is parallel to the floor and your shins are vertical.",
            "Squeeze your glutes hard at the top for a count of one or two.",
            "Lower your hips under control back toward the floor. Keep the movement in your hips and don't arch your lower back.",
        ],
        'Glute Bridge': [
            "Lie on your back with your knees bent and your feet flat on the floor, about hip-width apart.",
            f"Brace your core{(' and set ' + tool + ' across your hips') if eq != 'Bodyweight' else ''}.",
            "Push your hips up toward the ceiling by squeezing your glutes until your body forms a straight line from knees to shoulders.",
            "Hold the top for a second and feel your glutes contract.",
            "Lower your hips slowly back to the floor. Keep your ribs down and don't arch your lower back.",
        ],
        'Abduction': [
            f"Set up so {tool} pushes your knees or legs inward, with your hips stacked and your core braced.",
            "Push your knees or legs apart against the resistance, opening from the hips.",
            "Pause at the end of the range and squeeze your outer glutes.",
            "Bring your legs back together slowly under control, keeping the tension on your glutes the whole time.",
        ],
        'Adduction': [
            f"Set up with {tool} pulling your legs apart, hips stacked and core braced.",
            "Squeeze your legs together against the resistance, driving from your inner thighs.",
            "Pause for a moment at the point of full squeeze.",
            "Let your legs travel back apart slowly under control. Keep the movement smooth and don't let it snap back.",
        ],
        'Hip Extension': [
            "Set up on all fours, on a bench, or in the machine with your spine in a neutral, flat position and your core braced.",
            f"Take up the slack on {tool} so there is tension before you start.",
            "Drive your leg back and up using your glute, stopping when your hip is fully extended — don't over-arch your lower back.",
            "Squeeze your glute hard at the top for a count of one.",
            "Lower your leg slowly under control back to the start, keeping your spine neutral throughout.",
        ],
        'Activation': [
            "Lie on your side or set up on all fours with your hips stacked and your core braced.",
            "Keeping your feet or knees together where needed, lift slowly using your outer glute — keep your hips from rolling open or back.",
            "Pause at the top and feel the side of your glute working.",
            "Lower slowly and stay controlled. Move with the muscle, not with momentum.",
        ],
        'Lunge': [
            "Stand tall with your feet hip-width apart and your core braced.",
            "Step into a long stride and lower straight down until both knees are bent to about ninety degrees, keeping your front shin fairly upright.",
            "Drive through your front heel and squeeze your glute to stand back up.",
            "Keep your torso tall and your front knee tracking over your toes throughout.",
        ],
        'Step-up': [
            "Stand facing a sturdy box or bench at about knee height, feet hip-width apart and core braced.",
            "Place one whole foot flat on the box.",
            "Drive through that heel and squeeze your glute to stand all the way up on top of the box.",
            "Step back down slowly under control and repeat, keeping your chest up and your knee tracking over your toes.",
        ],
        'Squat': [
            "Stand with your feet about shoulder-width apart, toes turned out slightly, and your core braced.",
            "Push your hips back and bend your knees to lower down, keeping your chest up and your knees tracking over your toes.",
            "Descend until your thighs are at least parallel to the floor, keeping your weight through your mid-foot and heels.",
            "Drive through your heels and squeeze your glutes to stand back up tall. Don't let your knees cave in or your lower back round.",
        ],
        'Stretch/Mobility': [
            "Move into the stretch slowly until you feel a gentle pull through your glutes, hips or inner thighs.",
            "Hold the position and breathe deeply, letting the muscles lengthen with each exhale.",
            "Ease out of the stretch under control and repeat on the other side if needed.",
        ],
    }
    return STEPS.get(mv, [
        "Set up with a neutral, flat spine and your core braced.",
        "Drive your hips through and squeeze your glutes through the working range.",
        "Lower slowly under control and repeat. Don't arch your lower back.",
    ])


def folder_for(gender, mv):
    cat = CATEGORY.get(mv, 'Hip Extension & Kickbacks')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 - already organized.'); return
    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['category'] = CATEGORY.get(r['movement'], 'Hip Extension & Kickbacks')
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'])
        r['flag'] = ''
        r['folder'] = folder_for(r['gender'], r['movement'])
        recs.append(r)
    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Primary Muscle',
            'Secondary Muscles', 'Description', 'Step-by-step Instructions (TTS)', 'Variants', 'ID', 'Video File(s)', 'Folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r['name'], r['equipment'], r['gender'], r['level'], r['movement'], r['primary'],
                   r['secondary'], r['description'], '\n'.join(f"{i+1}. {s}" for i, s in enumerate(r['instructions'])),
                   r['variants'], r['id'], r['files'], r['folder']])
    ARIAL = 'Arial'
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'), 'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([38, 18, 9, 13, 20, 34, 38, 56, 70, 9, 11, 50, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 6, 7, 8, 9, 12, 13))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises', len(catalog)]); sm.append([])
    for t, k in [('By equipment', 'equipment'), ('By gender', 'gender'), ('By level', 'level'), ('By movement type', 'movement')]:
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    sm.column_dimensions['A'].width = 38; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\nSRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\nprint('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
