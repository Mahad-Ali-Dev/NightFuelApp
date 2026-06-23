"""
Core & Abs (Waist) catalog builder — same template as back/chest/biceps but with
WAIST / CORE / ABS exercise science (rectus abdominis, obliques, transverse
abdominis, hip flexors), and a step-by-step `instructions` array per exercise
written to be READ ALOUD by the app's speaker/TTS feature.
Run from this folder: python build_waist_catalog.py
"""
import os, re, json, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Core_And_Abs_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "waist_catalog.json")
TITLE = "Core & Abs"


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)
    name_part = s.split('_')[0]
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender, 'bodypart': 'Waist'}


def equipment(n):
    l = n.lower()
    if 'smith' in l: return 'Smith Machine'
    if 'landmine' in l: return 'T-Bar / Landmine'
    if 'barbell' in l: return 'Barbell'
    if 'dumbbell' in l: return 'Dumbbell'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'medicine ball' in l or 'medicine-ball' in l: return 'Medicine Ball'
    if 'cable' in l: return 'Cable'
    if 'ab roller' in l or 'ab wheel' in l or 'wheel rollout' in l or 'ab-roller' in l: return 'Ab Wheel'
    if 'sledgehammer' in l or 'battling rope' in l or 'battle rope' in l: return 'Conditioning Tool'
    if 'lever' in l or 'machine' in l or 'plate-loaded' in l or 'glute ham developer' in l or 'captain' in l: return 'Machine (Lever)'
    if re.search(r'\bband\b', l) or 'resistance band' in l or 'pallof' in l: return 'Resistance Band'
    if 'suspension' in l or 'trx' in l or 'suspender' in l or 'ring' in l: return 'Suspension (TRX)'
    if 'stability ball' in l or 'bosu' in l or 'swiss ball' in l: return 'Stability Ball'
    if 'hanging' in l or 'toes to bar' in l or 'parallel bar' in l or 'pull up' in l or 'pull-up' in l or 'pullup' in l or 'chin' in l: return 'Bodyweight (Bar)'
    if 'stick' in l: return 'Mobility Stick'
    return 'Bodyweight'


def movement(n):
    l = n.lower()
    # Stretch / mobility / self-myofascial (foam-rolling). `\broll\b` catches
    # "Roll Over" foam/spinal rolls but NOT "Rollout" (one word = no boundary =
    # the ab-wheel rollout, a CORE STABILITY move that must stay out of stretch).
    if any(k in l for k in ['stretch', 'mobility', 'foam']) or \
       (re.search(r'\broll\b', l) and 'rollout' not in l):
        return 'Stretch/Mobility'
    # Obliques / rotation / lateral flexion.
    if any(k in l for k in ['twist', 'russian', 'side bend', 'side-bend', 'woodchop', 'wood chop',
                            'wood-chop', 'oblique', 'bicycle', 'bycicle', 'windmill', 'pallof',
                            'rotation', 'rotary', 'side crunch', 'side plank', 'elbow to knee',
                            'side bridge', 'torso twist', 'obliques slides', 'crab twist']):
        return 'Oblique'
    # Anti-movement core stability (planks, holds, hollow, dead bug, bird dog, rollouts).
    if any(k in l for k in ['plank', 'hold', 'hollow', 'dead bug', 'dead-bug', 'bird dog', 'bird-dog',
                            'rollout', 'stability', 'planche', 'front lever', 'shoulder tap',
                            'bear', 'flutter', 'plank row', 'plank walk', 'walk up', 'pallof press']):
        return 'Core Stability'
    # Lower abs — leg / knee raises, reverse crunch, hanging work.
    if any(k in l for k in ['leg raise', 'leg-raise', 'knee raise', 'knee-raise', 'reverse crunch',
                            'leg lift', 'leg-lift', 'toes to bar', 'toes-to-bar', 'hanging',
                            'hip raise', 'hip lift', 'leg hip', 'pulse up', 'pulse-up',
                            'leg marches', 'flutter kick', 'knee raise', 'leg tuck', 'knee tuck',
                            'leg tap', 'tuck up', 'v tuck', 'knee in and out', 'knee raise']):
        return 'Lower Abs'
    # Upper abs — crunch / sit-up family.
    if any(k in l for k in ['crunch', 'sit up', 'sit-up', 'situp', 'curl-up', 'curl up',
                            'toe touch', 'heel touch', 'toe touching', 'v up', 'v-up', 'v sit',
                            'v-sit', 'jack knife', 'jack-knife', 'jackknife', 'starfish',
                            'pike', 'frog', 'groin']):
        return 'Upper Abs'
    return 'Upper Abs'


REGION = {  # movement -> training category
    'Upper Abs': 'Upper Abs', 'Lower Abs': 'Lower Abs', 'Oblique': 'Oblique',
    'Core Stability': 'Core Stability', 'Stretch/Mobility': 'Stretch',
}

ADV = ['hanging leg', 'hanging straight', 'hanging oblique', 'hanging toes', 'hanging half',
       'toes to bar', 'toes-to-bar', 'dragon flag', 'front lever', 'planche', 'windshield',
       'wiper', 'weighted', 'ab wheel rollout', 'wheel rollout', 'l-sit', 'l sit', 'l pull',
       'one arm', 'one-arm', 'ring ', 'pulse up', 'medicine ball slam', 'sledgehammer',
       'military press', 'vertical sit', 'vertical leg', 'captain']
BEG = ['stretch', 'mobility', 'foam', 'plank', 'crunch', 'dead bug', 'dead-bug', 'bird dog',
       'bird-dog', 'assisted', 'supported', 'wall', 'bridge', 'cat ', 'cobra', 'child',
       'beginner', 'kneeling', 'toe touch', 'heel touch', 'shoulder tap']


def level(n, mv):
    l = n.lower()
    if mv == 'Stretch/Mobility': return 'Beginner'
    if any(k in l for k in ADV): return 'Advanced'
    if any(k in l for k in BEG): return 'Beginner'
    return 'Intermediate'


def muscles(mv):
    return {
        'Upper Abs':      ('Rectus abdominis (upper)', 'Obliques, Hip flexors'),
        'Lower Abs':      ('Rectus abdominis (lower)', 'Hip flexors, Obliques'),
        'Oblique':        ('Obliques (internal & external)', 'Rectus abdominis, Transverse abdominis'),
        'Core Stability': ('Transverse abdominis & Rectus abdominis', 'Obliques, Erector spinae, Hip flexors'),
        'Stretch/Mobility': ('Abdominals & obliques (stretch)', 'Hip flexors, Erector spinae'),
    }.get(mv, ('Rectus abdominis (upper)', 'Obliques, Hip flexors'))


def describe(eq, mv):
    rl = {
        'Upper Abs': "a crunch-style movement that targets the upper abs — flexing the spine to pull the ribs toward the hips.",
        'Lower Abs': "a leg-raise style movement that emphasises the lower abs — lifting the legs while the lower back stays pinned.",
        'Oblique': "a rotational or side-bending movement that trains the obliques along the sides of the trunk.",
        'Core Stability': "an anti-movement core drill that braces the trunk and resists motion to build deep stability.",
        'Stretch/Mobility': "an abdominal stretch and mobility drill to lengthen the abs and obliques and open the spine.",
    }.get(mv, "a core movement for the abs.")
    # role strings already begin with their own article ("a"/"an"); just capitalise.
    sentence = rl[0].upper() + rl[1:]
    eqs = '' if eq == 'Bodyweight' else f" Performed with {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}."
    return sentence + eqs


def instructions(eq, mv, name):
    """Step-by-step how-to, written to be READ ALOUD (speaker/TTS). Returns a list.
    Numbers are spelled out. Always cue bracing the core and never yanking the neck."""
    l = name.lower()
    load = 'the cable' if eq == 'Cable' else 'the weight' if eq in ('Dumbbell', 'Kettlebell', 'Barbell', 'Medicine Ball') else 'the wheel' if eq == 'Ab Wheel' else 'the band'

    # Cable / hanging family — specific cueing wins over the generic bucket steps.
    if eq == 'Cable' or 'cable' in l:
        return [
            f"Set the pulley and take hold of {load}, then settle into a stable position with your core already braced.",
            "Set your hips and let the weight load your abs, keeping a slight tension on the line so it never goes slack.",
            "Move slowly and crunch or rotate against the resistance, exhaling as you reach the fully contracted position.",
            "Squeeze your abs hard for a second at the end of the range.",
            "Return under control to the start, keeping the tension the whole way. Brace your core and never yank with your neck or arms.",
        ]
    if 'hanging' in l or 'toes to bar' in l or 'toes-to-bar' in l:
        return [
            "Hang from the bar with your arms straight and pull your shoulders down to take the slack out.",
            "Brace your core and tilt your pelvis so your lower abs are switched on before you move.",
            "Raise your knees or legs under control, curling your pelvis up toward your ribs rather than just swinging.",
            "Exhale as you reach the top and pause for a moment without swinging.",
            "Lower slowly to a full hang, staying in control. Keep the movement smooth and never let momentum take over.",
        ]

    STEPS = {
        'Upper Abs': [
            "Lie on your back with your knees bent and your feet flat, and rest your hands lightly by your head or across your chest.",
            "Press your lower back gently into the floor and brace your abs.",
            "Curl your shoulder blades up off the floor by contracting your abs, keeping your chin off your chest and your neck relaxed.",
            "Pause for a moment at the top and squeeze your abs.",
            "Lower slowly back down with control. Brace your core throughout and never pull on your neck with your hands.",
        ],
        'Lower Abs': [
            "Lie on your back with your legs extended and your hands tucked just under your hips or by your sides.",
            "Press your lower back firmly into the floor so it stays flat the whole time.",
            "Raise your legs under control, lifting from your lower abs while keeping your back pinned down.",
            "Pause briefly at the top, then lower your legs slowly without letting your back arch up off the floor.",
            "Brace your core throughout and stop the range short of any point where your lower back lifts.",
        ],
        'Oblique': [
            "Set up in a stable position with your core braced and your spine long.",
            "Lean back slightly if you are seated, or fix your hips if you are standing, so the work stays in your sides.",
            "Rotate your torso, or bend toward one side, in a slow and controlled arc using your obliques.",
            "Squeeze the side of your waist at the end of the range, then return through the middle under control.",
            "Repeat smoothly to the other side. Brace your core, move deliberately, and never twist or yank with your neck.",
        ],
        'Core Stability': [
            "Set up on your forearms or hands with your body in one straight line from your head to your heels.",
            "Brace your abs and squeeze your glutes so your hips do not sag or pike up.",
            "Hold the position, keeping your neck long and your gaze down at the floor.",
            "Breathe steadily and keep everything tight for the full duration of the hold.",
            "Keep your core braced the entire time and never let your lower back drop or your neck strain.",
        ],
        'Stretch/Mobility': [
            "Move into the stretch slowly until you feel a gentle pull through your abs, sides, or lower back.",
            "Hold the position and breathe deeply, letting the muscles lengthen a little more with each exhale.",
            "Ease out of the stretch under control and repeat on the other side if needed. Keep your neck relaxed throughout.",
        ],
    }
    return STEPS.get(mv, [
        "Set up in a stable position with your core already braced.",
        "Move slowly through the range, contracting your abs as you go.",
        "Return under control and repeat. Brace your core and never yank with your neck.",
    ])


def folder_for(gender, mv):
    cat = {
        'Upper Abs': 'Upper Abs (Crunch)',
        'Lower Abs': 'Lower Abs (Leg Raise)',
        'Oblique': 'Obliques (Twist & Side)',
        'Core Stability': 'Core Stability (Plank)',
        'Stretch': 'Stretch & Mobility',
    }.get(REGION.get(mv, 'Upper Abs'), 'Other')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 - already organized.'); return
    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['region'] = REGION.get(r['movement'], 'Upper Abs')
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'], r['name'])
        r['flag'] = ''
        r['folder'] = folder_for(r['gender'], r['movement'])
        recs.append(r)
    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Primary Muscle',
            'Secondary Muscles', 'Description', 'Step-by-step Instructions (TTS)', 'Variants', 'ID', 'Video File(s)', 'Folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r['name'], r['equipment'], r['gender'], r['level'], r['movement'], r['primary'],
                   r['secondary'], r['description'], '\n'.join(f"{i+1}. {s}" for i, s in enumerate(r['instructions'])),
                   r['variants'], r['id'], r['files'], r['folder']])
    ARIAL = 'Arial'
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'), 'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([38, 18, 9, 13, 20, 34, 38, 56, 70, 9, 11, 50, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 6, 7, 8, 9, 12, 13))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises', len(catalog)]); sm.append([])
    for t, k in [('By equipment', 'equipment'), ('By gender', 'gender'), ('By level', 'level'), ('By movement type', 'movement')]:
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    sm.column_dimensions['A'].width = 38; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\nSRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\nprint('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
