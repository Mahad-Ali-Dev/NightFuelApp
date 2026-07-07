"""
Back catalog builder — same template as chest/biceps but with BACK exercise
science (lats / traps / rhomboids / erectors), and a NEW step-by-step
`instructions` array per exercise written to be READ ALOUD by the app's
speaker/TTS feature. Run from this folder: python build_back_catalog.py
"""
import os, re, json, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Back_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "back_catalog.json")
TITLE = "Back"


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)
    name_part = s.split('_')[0]
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender, 'bodypart': 'Back'}


def equipment(n):
    l = n.lower()
    if 'smith' in l: return 'Smith Machine'
    if 't-bar' in l or 't bar' in l or 'landmine' in l: return 'T-Bar / Landmine'
    if 'barbell' in l: return 'Barbell'
    if 'dumbbell' in l: return 'Dumbbell'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'cable' in l: return 'Cable'
    if 'machine' in l or 'lever' in l or 'pulldown machine' in l or 'plate-loaded' in l: return 'Machine (Lever)'
    if re.search(r'\bband\b', l) or 'resistance band' in l: return 'Resistance Band'
    if 'suspension' in l or 'trx' in l: return 'Suspension (TRX)'
    if 'pull-up' in l or 'pull up' in l or 'pullup' in l or 'chin' in l: return 'Bodyweight (Bar)'
    return 'Bodyweight'


def movement(n):
    l = n.lower()
    # Stretch / mobility / self-myofascial (foam-rolling). `\broll\b` catches
    # "Roll Upper Back" etc. but NOT "Rollout" (one word = no boundary = a core move).
    if any(k in l for k in ['stretch', 'mobility', 'opener', 'cat ', 'child', 'foam',
                            'roller', 'release', 'decompress', 'dead hang']) or re.search(r'\broll\b', l):
        return 'Stretch/Mobility'
    if 'shrug' in l: return 'Shrug'
    if 'pullover' in l: return 'Pullover'
    if 'straight' in l and ('pushdown' in l or 'pulldown' in l): return 'Straight-Arm Pulldown'
    if 'pushdown' in l: return 'Straight-Arm Pulldown'
    if any(k in l for k in ['hyperextension', 'back extension', 'good morning', 'good-morning', 'superman', 'reverse hyper']):
        return 'Lower-Back Extension'
    if 'deadlift' in l or 'rack pull' in l: return 'Deadlift'
    if 'pulldown' in l or 'pull-down' in l: return 'Pulldown'
    if 'pull-up' in l or 'pull up' in l or 'pullup' in l or 'chin' in l: return 'Pull-up'
    if 'row' in l: return 'Row'
    return 'Row'


REGION = {  # movement -> training category
    'Pulldown': 'Vertical Pull', 'Pull-up': 'Vertical Pull', 'Straight-Arm Pulldown': 'Vertical Pull',
    'Row': 'Horizontal Pull', 'Shrug': 'Traps', 'Pullover': 'Pullover',
    'Deadlift': 'Lower Back', 'Lower-Back Extension': 'Lower Back', 'Stretch/Mobility': 'Stretch',
}

ADV = ['one-arm', 'one arm', 'single-arm', 'single arm', 'deficit', 'snatch grip', 'pendlay',
       'archer', 'pull-up', 'pull up', 'pullup', 'muscle-up', 'l-sit', 'weighted', 'around the world', 'kroc']
BEG = ['machine', 'lever', 'assisted', 'supported', 'chest supported', 'seal', 'cable', 'band',
       'stretch', 'mobility', 'hyperextension', 'superman', 'shrug', 'wall']


def level(n, mv):
    l = n.lower()
    if mv == 'Stretch/Mobility': return 'Beginner'
    if mv == 'Deadlift': return 'Advanced' if any(k in l for k in ['deficit', 'snatch', 'one', 'single', 'stiff']) else 'Intermediate'
    if any(k in l for k in ADV): return 'Advanced'
    if any(k in l for k in BEG): return 'Beginner'
    return 'Intermediate'


def muscles(mv):
    return {
        'Row':                  ('Latissimus dorsi, Rhomboids, Middle trapezius', 'Posterior deltoid, Biceps brachii, Erector spinae'),
        'Pulldown':             ('Latissimus dorsi', 'Teres major, Rhomboids, Biceps brachii, Posterior deltoid'),
        'Pull-up':              ('Latissimus dorsi', 'Teres major, Rhomboids, Biceps brachii, Core'),
        'Straight-Arm Pulldown':('Latissimus dorsi', 'Teres major, Long head triceps, Posterior deltoid'),
        'Shrug':                ('Upper trapezius', 'Levator scapulae, Forearm flexors (grip)'),
        'Pullover':             ('Latissimus dorsi', 'Teres major, Serratus anterior, Pectoralis major, Long head triceps'),
        'Deadlift':             ('Erector spinae, Gluteus maximus, Hamstrings', 'Latissimus dorsi, Trapezius, Forearm flexors'),
        'Lower-Back Extension': ('Erector spinae', 'Gluteus maximus, Hamstrings'),
        'Stretch/Mobility':     ('Latissimus dorsi & spinal extensors (stretch)', 'Rhomboids, Trapezius'),
    }.get(mv, ('Latissimus dorsi', 'Biceps brachii, Rhomboids'))


def describe(eq, mv):
    rl = {
        'Row': "a horizontal pull that builds mid-back thickness — lats, rhomboids and mid-traps — with the rear delts and biceps assisting.",
        'Pulldown': "a vertical pull that builds lat WIDTH, pulling a bar down to the upper chest.",
        'Pull-up': "a bodyweight vertical pull — one of the best lat and upper-back builders.",
        'Straight-Arm Pulldown': "an isolation pull with locked elbows that targets the lats directly without much biceps help.",
        'Shrug': "an isolation lift for the upper traps — elevating the shoulders against load.",
        'Pullover': "an overhead pull that stretches and works the lats and serratus across a long range.",
        'Deadlift': "a full posterior-chain hinge — erectors, glutes and hamstrings — with the whole back bracing.",
        'Lower-Back Extension': "a hip-hinge that strengthens the lower-back erectors with the glutes and hamstrings.",
        'Stretch/Mobility': "a back stretch/mobility drill to lengthen the lats and decompress the spine.",
    }.get(mv, "a pulling movement for the back.")
    # rl already begins with its own article ("a"/"an") — just capitalize it
    # (no extra article, which would double to "An a horizontal pull…").
    eqs = '' if eq == 'Bodyweight' else f" Performed with {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}."
    return rl[0].upper() + rl[1:] + eqs


def instructions(eq, mv):
    """Step-by-step how-to, written to be READ ALOUD (speaker/TTS). Returns a list."""
    tool = 'the bar' if 'Bar' in eq or eq in ('Barbell', 'Smith Machine', 'T-Bar / Landmine') else \
           'the handle' if eq == 'Cable' else 'the weight' if eq in ('Dumbbell', 'Kettlebell') else 'the handles'
    STEPS = {
        'Row': [
            f"Stand with your feet hip-width apart and hinge forward at the hips until your torso is at roughly forty-five degrees, keeping your back flat.",
            f"Let {tool} hang straight down with your arms fully extended and your shoulders relaxed.",
            f"Pull {tool} toward your lower ribs, driving your elbows back and squeezing your shoulder blades together.",
            "Pause for a moment at the top and feel your mid-back contract.",
            f"Lower {tool} slowly until your arms are fully straight, keeping your torso steady. Lead with your elbows, not your hands, and don't yank with your lower back.",
        ],
        'Pulldown': [
            f"Sit down and grip {tool} a little wider than shoulder-width, then set your shoulders down and back.",
            "Lean back very slightly and lift your chest toward the bar.",
            f"Pull {tool} down to your upper chest by driving your elbows down toward your sides.",
            "Squeeze your lats hard at the bottom for a second.",
            f"Let {tool} rise slowly back to a full overhead stretch, staying in control the whole way.",
        ],
        'Pull-up': [
            "Grip the bar slightly wider than shoulder-width with your palms facing away.",
            "Hang with your arms straight, then pull your shoulders down and brace your core.",
            "Pull your chest toward the bar by driving your elbows down and back until your chin clears it.",
            "Squeeze your lats at the top.",
            "Lower yourself slowly to a full hang. Avoid swinging or kipping if you're training for strength.",
        ],
        'Straight-Arm Pulldown': [
            f"Stand facing the cable, grip {tool} with arms extended in front of you and a slight bend in the elbows.",
            "Hinge slightly at the hips and brace your core.",
            f"Keeping your arms nearly straight, pull {tool} down in an arc to your thighs using only your lats.",
            "Squeeze your lats at the bottom.",
            f"Let {tool} rise back overhead under control, feeling the stretch.",
        ],
        'Shrug': [
            f"Stand tall holding {tool} at arm's length by your sides, shoulders relaxed.",
            "Lift your shoulders straight up toward your ears as high as you can.",
            "Hold the squeeze at the very top for a moment.",
            "Lower slowly back to the start. Keep your arms straight and don't roll your shoulders.",
        ],
        'Pullover': [
            f"Lie back on the bench and hold {tool} above your chest with a slight bend in your elbows.",
            f"Lower {tool} back behind your head until you feel a deep stretch across your lats and chest.",
            f"Pull {tool} back over your chest using your lats, keeping your elbows fixed.",
            "Keep your ribs down and your core braced throughout.",
        ],
        'Deadlift': [
            f"Set up with {tool} over the middle of your feet, feet hip-width apart.",
            "Hinge at the hips and grip the bar with a flat back, chest up and shoulders just in front of the bar.",
            "Take a breath, brace your core, then drive through your heels and stand tall, keeping the bar close to your body.",
            "Squeeze your glutes at the top — don't lean back.",
            "Push your hips back and lower the bar under control along the same path. Keep your spine neutral the whole time — never round your lower back.",
        ],
        'Lower-Back Extension': [
            "Position your hips on the pad with your feet anchored under the supports.",
            "Cross your arms or hold a light weight at your chest, with your body in a straight line.",
            "Hinge at the hips and lower your torso toward the floor with a flat back.",
            "Squeeze your glutes and lower back to raise your torso back to straight — don't over-arch at the top.",
        ],
        'Stretch/Mobility': [
            "Move into the stretch slowly until you feel a gentle pull through your back.",
            "Hold the position and breathe deeply, letting the muscles lengthen with each exhale.",
            "Ease out of the stretch under control and repeat on the other side if needed.",
        ],
    }
    return STEPS.get(mv, [
        f"Set up with a flat back and a firm grip on {tool}.",
        f"Pull {tool} toward your body, squeezing your back muscles.",
        f"Lower {tool} slowly under control and repeat.",
    ])


def folder_for(gender, mv):
    cat = {
        'Vertical Pull': 'Vertical Pull (Lats & Width)',
        'Horizontal Pull': 'Horizontal Pull (Rows & Mid-Back)',
        'Traps': 'Traps (Shrugs)',
        'Pullover': 'Pullover (Lats & Serratus)',
        'Lower Back': 'Lower Back (Deadlift & Extension)',
        'Stretch': 'Stretch & Mobility',
    }.get(REGION.get(mv, 'Horizontal Pull'), 'Other')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 - already organized.'); return
    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['region'] = REGION.get(r['movement'], 'Horizontal Pull')
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'])
        r['flag'] = ''
        r['folder'] = folder_for(r['gender'], r['movement'])
        recs.append(r)
    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Primary Muscle',
            'Secondary Muscles', 'Description', 'Step-by-step Instructions (TTS)', 'Variants', 'ID', 'Video File(s)', 'Folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r['name'], r['equipment'], r['gender'], r['level'], r['movement'], r['primary'],
                   r['secondary'], r['description'], '\n'.join(f"{i+1}. {s}" for i, s in enumerate(r['instructions'])),
                   r['variants'], r['id'], r['files'], r['folder']])
    ARIAL = 'Arial'
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'), 'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([38, 18, 9, 13, 20, 34, 38, 56, 70, 9, 11, 50, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 6, 7, 8, 9, 12, 13))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises', len(catalog)]); sm.append([])
    for t, k in [('By equipment', 'equipment'), ('By gender', 'gender'), ('By level', 'level'), ('By movement type', 'movement')]:
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    sm.column_dimensions['A'].width = 38; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\nSRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\nprint('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
