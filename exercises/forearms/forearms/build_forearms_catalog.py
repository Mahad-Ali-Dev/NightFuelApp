"""
Forearms catalog builder — same template as back/chest/biceps but with FOREARMS
exercise science (wrist flexors / extensors / brachioradialis / grip), and a
step-by-step `instructions` array per exercise written to be READ ALOUD by the
app's speaker/TTS feature. Run from this folder: python build_forearms_catalog.py
"""
import os, re, json, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Forearms_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "forearms_catalog.json")
TITLE = "Forearms"


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)
    name_part = s.split('_')[0]
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender, 'bodypart': 'Forearms'}


def equipment(n):
    l = n.lower()
    if 'smith' in l: return 'Smith Machine'
    if 'ez' in l.split() or 'ez-' in l or 'ez ' in l: return 'EZ Bar'
    if 'barbell' in l: return 'Barbell'
    if 'dumbbell' in l: return 'Dumbbell'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'cable' in l: return 'Cable'
    if 'machine' in l or 'lever' in l or 'plate-loaded' in l: return 'Machine (Lever)'
    if re.search(r'\bband\b', l) or 'resistance band' in l: return 'Resistance Band'
    if 'roller' in l: return 'Wrist Roller'
    if 'plate' in l: return 'Weight Plate'
    return 'Bodyweight'


def movement(n):
    l = n.lower()
    # Stretch / mobility / self-myofascial (foam-rolling). `\broll\b` catches
    # "Roll" as a word but NOT "Roller"/"Rollout" (no word boundary).
    if any(k in l for k in ['stretch', 'mobility', 'foam']) or re.search(r'\broll\b', l):
        return 'Stretch/Mobility'
    if 'reverse curl' in l:
        return 'Reverse Curl'
    if ('reverse' in l and 'wrist' in l) or 'wrist extension' in l:
        return 'Wrist Extension'
    if 'wrist curl' in l or 'wrist flexion' in l:
        return 'Wrist Flexion'
    if 'roller' in l:
        return 'Wrist Roller'
    if any(k in l for k in ['farmer', 'carry', 'hold', 'grip', 'hang', 'pinch']):
        return 'Grip & Carry'
    if any(k in l for k in ['pronation', 'supination', 'twist']):
        return 'Rotation'
    if 'wrist' in l:
        return 'Wrist Flexion'
    return 'Wrist Flexion'


REGION = {  # movement -> training category (folder, NO "/" — use "&")
    'Wrist Flexion': 'Wrist Flexion',
    'Wrist Extension': 'Wrist Extension',
    'Reverse Curl': 'Reverse (Brachioradialis)',
    'Wrist Roller': 'Grip & Carry',
    'Grip & Carry': 'Grip & Carry',
    'Rotation': 'Pronation & Supination',
    'Stretch/Mobility': 'Stretch & Mobility',
}

# heavy / unilateral / thick-bar grip work is the most demanding
ADV = ['farmer', 'one-arm', 'one arm', 'single-arm', 'single arm', 'pinch', 'thick-bar', 'thick bar', 'heavy']
BEG = ['machine', 'lever', 'cable', 'band', 'seated wrist']


def level(n, mv):
    l = n.lower()
    if mv == 'Stretch/Mobility':
        return 'Beginner'
    if any(k in l for k in ADV):
        return 'Advanced'
    if any(k in l for k in BEG) or 'seated' in l and 'wrist' in l:
        return 'Beginner'
    if 'machine' in l or 'cable' in l or re.search(r'\bband\b', l):
        return 'Beginner'
    return 'Intermediate'


def muscles(mv):
    return {
        'Wrist Flexion':   ('Wrist flexors (flexor carpi radialis & ulnaris)', 'Finger flexors'),
        'Wrist Extension': ('Wrist extensors (extensor carpi radialis & ulnaris)', 'Finger extensors'),
        'Reverse Curl':    ('Brachioradialis', 'Wrist extensors, Biceps brachii'),
        'Wrist Roller':    ('Wrist flexors & extensors', 'Finger flexors (grip)'),
        'Grip & Carry':    ('Finger flexors (flexor digitorum)', 'Wrist flexors, Forearm stabilizers'),
        'Rotation':        ('Pronator teres & Supinator', 'Brachioradialis'),
        'Stretch/Mobility':('Forearm flexors & extensors (stretch)', 'Wrist'),
    }.get(mv, ('Wrist flexors (flexor carpi radialis & ulnaris)', 'Finger flexors'))


def describe(eq, mv):
    rl = {
        'Wrist Flexion': "a wrist-curl that targets the forearm flexors on the underside of the forearm, building grip and forearm thickness.",
        'Wrist Extension': "a reverse wrist-curl that targets the forearm extensors on top of the forearm for balanced, healthy wrists.",
        'Reverse Curl': "a palms-down curl that builds the brachioradialis and forearm extensors, adding width to the upper forearm.",
        'Wrist Roller': "a continuous winding drill that torches the wrist flexors and extensors and builds serious grip endurance.",
        'Grip & Carry': "a loaded carry or hold that hammers the finger flexors and grip while the whole forearm braces.",
        'Rotation': "a forearm rotation drill working the pronators and supinators that turn the wrist.",
        'Stretch/Mobility': "a forearm stretch and mobility drill to lengthen the wrist flexors and extensors.",
    }.get(mv, "a forearm movement that builds wrist and grip strength.")
    art = 'An ' if rl[0] in 'AEIOUaeiou' else 'A '
    eqs = '' if eq == 'Bodyweight' else f" Performed with {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}."
    return art + rl + eqs


def instructions(eq, mv):
    """Step-by-step how-to, written to be READ ALOUD (speaker/TTS). Returns a list."""
    tool = 'the bar' if 'Bar' in eq or eq in ('Barbell', 'Smith Machine') else \
           'the handle' if eq == 'Cable' else 'the weight' if eq in ('Dumbbell', 'Kettlebell', 'Weight Plate') else \
           'the roller' if eq == 'Wrist Roller' else 'the weight'
    STEPS = {
        'Wrist Flexion': [
            f"Sit down and rest your forearms on your thighs or a bench, palms facing up, with {tool} held in your hands and your wrists just past your knees.",
            f"Let {tool} roll down toward your fingertips, opening your hands and feeling the stretch through your forearms.",
            "Curl your hands back up by flexing only your wrists, keeping your forearms flat and still.",
            "Squeeze hard at the top for a moment.",
            f"Lower {tool} slowly under control through a full range and repeat. Move only at the wrists — keep the rest of your arms quiet.",
        ],
        'Wrist Extension': [
            f"Sit down and rest your forearms on your thighs or a bench, palms facing DOWN, with {tool} held and your wrists just past your knees.",
            f"Let your hands drop so {tool} lowers toward the floor, feeling a stretch on top of your forearms.",
            "Raise the backs of your hands as high as you can by extending your wrists, keeping your forearms pinned down.",
            "Hold the squeeze at the top for a second.",
            f"Lower {tool} slowly under control and repeat. Use a light load and a full range — this is a small muscle.",
        ],
        'Reverse Curl': [
            f"Stand tall and grip {tool} with your palms facing down, hands about shoulder-width apart and arms straight.",
            "Keep your upper arms pinned to your sides and your wrists firm and straight.",
            "Curl the weight up by bending only at the elbows, leading with the backs of your hands.",
            "Squeeze your forearms and brachioradialis at the top.",
            f"Lower {tool} slowly all the way down under control and repeat. Don't swing — keep the upper arms still.",
        ],
        'Wrist Roller': [
            f"Hold {tool} out in front of you at about shoulder height with both hands, arms straight, and the cord hanging with the weight at the bottom.",
            "Brace your core and keep your arms extended.",
            "Wind the weight up by rolling the bar with one wrist at a time, hand over hand, until the weight reaches the top.",
            "Then reverse the motion and lower the weight slowly with the same continuous wrist action — don't let it drop.",
            "Keep the movement smooth and constant. Your forearms will burn — that's the point.",
        ],
        'Grip & Carry': [
            f"Stand tall holding {tool} at your sides with a firm, full grip and your shoulders back.",
            "Brace your core, pull your shoulders down, and stand as tall as you can.",
            "Walk forward with short, controlled steps, keeping the weight steady and your posture upright — or simply hold the position and grip hard if you're not walking.",
            "Keep crushing your grip the whole time and breathe steadily.",
            "Set the weight down under control when your grip is nearly spent. Never let it slip.",
        ],
        'Rotation': [
            f"Rest your forearm along your thigh or a bench with your hand off the end, holding {tool} so it sits to one side.",
            "Keep your forearm still and your wrist in line with it.",
            "Rotate your forearm to turn your palm up and then down through a full, controlled range.",
            "Pause briefly at each end and feel the muscles that turn the wrist working.",
            f"Continue slowly under control and repeat. Move only the forearm — keep the elbow steady.",
        ],
        'Stretch/Mobility': [
            "Extend one arm in front of you and use your other hand to gently bend the wrist until you feel a stretch through the forearm.",
            "Hold the position and breathe deeply, letting the muscles lengthen with each exhale.",
            "Ease out of the stretch under control, then switch directions to stretch both the flexors and the extensors.",
        ],
    }
    return STEPS.get(mv, [
        f"Set up with your forearm supported and a firm grip on {tool}.",
        f"Move only at the wrist, curling {tool} through a full range.",
        f"Lower {tool} slowly under control and repeat.",
    ])


def folder_for(gender, mv):
    cat = REGION.get(mv, 'Wrist Flexion')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 - already organized.'); return
    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['region'] = REGION.get(r['movement'], 'Wrist Flexion')
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'])
        r['flag'] = ''
        r['folder'] = folder_for(r['gender'], r['movement'])
        recs.append(r)
    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Primary Muscle',
            'Secondary Muscles', 'Description', 'Step-by-step Instructions (TTS)', 'Variants', 'ID', 'Video File(s)', 'Folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r['name'], r['equipment'], r['gender'], r['level'], r['movement'], r['primary'],
                   r['secondary'], r['description'], '\n'.join(f"{i+1}. {s}" for i, s in enumerate(r['instructions'])),
                   r['variants'], r['id'], r['files'], r['folder']])
    ARIAL = 'Arial'
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'), 'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([38, 18, 9, 13, 20, 34, 38, 56, 70, 9, 11, 50, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 6, 7, 8, 9, 12, 13))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises', len(catalog)]); sm.append([])
    for t, k in [('By equipment', 'equipment'), ('By gender', 'gender'), ('By level', 'level'), ('By movement type', 'movement')]:
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    sm.column_dimensions['A'].width = 38; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\nSRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\nprint('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
