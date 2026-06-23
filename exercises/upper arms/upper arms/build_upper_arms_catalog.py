"""
Upper Arms catalog builder — same proven template as back/biceps/triceps, but
this folder holds COMBINED biceps + triceps demos (curls AND extensions /
pushdowns / dips / presses). It merges the two sciences: it routes each clip to
a BICEPS or TRICEPS movement, picks the matching primary/secondary muscles, and
writes a step-by-step `instructions` array per exercise meant to be READ ALOUD
by the app's speaker/TTS feature.  Run from this folder:
    python build_upper_arms_catalog.py
"""
import os, re, json, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Upper_Arms_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "upper_arms_catalog.json")
TITLE = "Upper Arms"


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*-\s*Copy.*$', '', s, flags=re.I)   # " - Copy", " - Copy - Copy"
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)               # trailing (1)(2) dupes
    name_part = s.split('_')[0]
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = re.sub(r'\(VERSIO$', '', name, flags=re.I)   # truncated "(VERSION-2)" tails
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender, 'bodypart': 'Upper Arms'}


def equipment(n):
    l = n.lower()
    if 'smith' in l: return 'Smith Machine'
    if 'ez' in l.split() or 'ez-bar' in l or 'ez bar' in l or 'sz-bar' in l or 'sz bar' in l: return 'EZ / SZ Bar'
    if 'barbell' in l or 'olympic' in l: return 'Barbell'
    if 'dumbbell' in l: return 'Dumbbell'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'cable' in l: return 'Cable'
    if 'machine' in l or 'lever' in l or 'plate-loaded' in l: return 'Machine (Lever)'
    if re.search(r'\bband\b', l) or 'resistance band' in l: return 'Resistance Band'
    if 'suspension' in l or 'suspender' in l or 'trx' in l or 'ring' in l: return 'Suspension (TRX)'
    if 'towel' in l or 'bottle' in l: return 'Household / Improvised'
    return 'Bodyweight'


# --- the two halves of the movement vocabulary ---------------------------------
BICEPS_MOVES = {'Curl', 'Incline Curl', 'Drag Curl', 'Preacher Curl', 'Spider Curl',
                'Concentration Curl', 'Hammer Curl', 'Reverse Curl', 'Zottman Curl'}
TRICEPS_MOVES = {'Pushdown', 'Overhead Extension', 'Lying Extension', 'Close-Grip Press',
                 'Dip', 'Kickback'}


def movement(n):
    l = n.lower()
    # 1) Stretch / mobility / self-myofascial (foam-rolling) FIRST. `\broll\b`
    #    catches "Roll ..." but NOT "Rollout" (one word = no boundary = a core move).
    if any(k in l for k in ['stretch', 'mobility', 'opener', 'foam', 'roller', 'release',
                            'decompress', 'dead hang', 'downward', 'doorway']) or re.search(r'\broll\b', l):
        return 'Stretch/Mobility'

    # 2) TRICEPS keywords
    if 'kickback' in l: return 'Kickback'
    if 'skull' in l or 'lying triceps' in l or 'french' in l or 'tate' in l: return 'Lying Extension'
    if 'pushdown' in l or 'press-down' in l or 'pressdown' in l: return 'Pushdown'
    if ('close' in l and 'grip' in l) or 'jm ' in l or 'diamond' in l or ('narrow' in l and 'press' in l) or 'bench press' in l:
        return 'Close-Grip Press'
    if 'dip' in l: return 'Dip'
    if 'overhead' in l or 'triceps extension' in l or 'tricep extension' in l or ('extension' in l and 'leg' not in l):
        return 'Overhead Extension'

    # 3) BICEPS keywords
    if 'zottman' in l: return 'Zottman Curl'
    if 'reverse' in l and 'curl' in l: return 'Reverse Curl'
    if 'hammer' in l: return 'Hammer Curl'
    if 'preacher' in l: return 'Preacher Curl'
    if 'spider' in l or 'prone incline' in l: return 'Spider Curl'
    if 'concentration' in l: return 'Concentration Curl'
    if 'drag' in l: return 'Drag Curl'
    if 'incline' in l and 'curl' in l: return 'Incline Curl'
    if 'curl' in l: return 'Curl'

    # 4) Leftovers: route to whichever side the remaining words imply.
    #    Triceps signals (presses/drives/dips/kickback-style/inverted press), else
    #    biceps ("21s" is a classic biceps curl protocol).
    if 'tricep' in l or 'rear drive' in l or 'press' in l or 'lockout' in l or 'handstand' in l:
        return 'Overhead Extension'
    return 'Curl'


def is_biceps(mv):
    return mv in BICEPS_MOVES


# head-emphasis sub-region (used only for the biceps folder split)
def biceps_region(mv):
    return {
        'Incline Curl': 'Long Head', 'Drag Curl': 'Long Head', 'Curl': 'Overall',
        'Spider Curl': 'Curls', 'Preacher Curl': 'Short Head', 'Concentration Curl': 'Short Head',
        'Hammer Curl': 'Brachialis', 'Reverse Curl': 'Brachialis', 'Zottman Curl': 'Brachialis',
    }.get(mv, 'Overall')


ADV = ['one-arm', 'one arm', 'single-arm', 'single arm', 'unilateral', 'deficit',
       'weighted dip', 'weighted-dip', 'weighted three', 'zottman', 'drag', 'concentration',
       'ring', 'handstand', 'side-lying', 'side lying']
BEG = ['assisted', 'machine', 'lever', 'cable', 'band', 'resistance band', 'bench dip',
       'stretch', 'mobility', 'wall']


def level(n, mv):
    l = n.lower()
    if mv == 'Stretch/Mobility': return 'Beginner'
    if any(k in l for k in ADV): return 'Advanced'
    if any(k in l for k in BEG): return 'Beginner'
    return 'Intermediate'


# --- muscle maps: copied verbatim from the biceps + triceps builders -----------
BICEPS_MUSCLES = {
    'Curl':               ('Biceps brachii', 'Brachialis, Brachioradialis'),
    'Incline Curl':       ('Biceps brachii (long head)', 'Brachialis'),
    'Drag Curl':          ('Biceps brachii (long head)', 'Brachialis, Posterior deltoid'),
    'Preacher Curl':      ('Biceps brachii (short head)', 'Brachialis'),
    'Spider Curl':        ('Biceps brachii (short head)', 'Brachialis'),
    'Concentration Curl': ('Biceps brachii (short head, peak)', 'Brachialis'),
    'Hammer Curl':        ('Brachialis & Brachioradialis', 'Biceps brachii'),
    'Reverse Curl':       ('Brachioradialis & wrist extensors', 'Biceps brachii, Brachialis'),
    'Zottman Curl':       ('Biceps brachii & Brachioradialis', 'Brachialis, Wrist extensors'),
}
TRICEPS_MUSCLES = {
    'Overhead Extension':  ('Triceps brachii (long head)', 'Anconeus'),
    'Pushdown':            ('Triceps brachii (lateral head)', 'Anconeus'),
    'Lying Extension':     ('Triceps brachii (long & lateral head)', 'Anconeus'),
    'Kickback':            ('Triceps brachii (lateral & long head)', 'Anconeus'),
    'Close-Grip Press':    ('Triceps brachii', 'Pectoralis major, Anterior deltoid'),
    'Dip':                 ('Triceps brachii', 'Pectoralis major (lower), Anterior deltoid'),
}


def muscles(mv):
    if mv == 'Stretch/Mobility':
        return ('Biceps & Triceps brachii (stretch)', 'Brachialis, Posterior deltoid, Anterior forearm')
    if is_biceps(mv):
        return BICEPS_MUSCLES.get(mv, ('Biceps brachii', 'Brachialis, Brachioradialis'))
    return TRICEPS_MUSCLES.get(mv, ('Triceps brachii', 'Anconeus'))


def describe(eq, mv):
    """One clean grammatical sentence with a SINGLE leading article."""
    if is_biceps(mv):
        rl = {
            'Curl': "biceps curl that lifts the load from a full stretch to a hard squeeze with the palms up, keeping the elbows pinned to the sides so the biceps do the work.",
            'Incline Curl': "incline curl performed with the arms hanging behind the torso to bias the long (outer) head of the biceps.",
            'Drag Curl': "drag curl that pulls the load up close to the body with the elbows tracking back, keeping tension on the long head through the top.",
            'Preacher Curl': "preacher curl with the upper arms fixed on the pad, curling from a deep stretch to kill momentum and load the short (inner) head.",
            'Spider Curl': "spider curl done chest-down on an incline with the arms hanging straight for constant short-head tension.",
            'Concentration Curl': "concentration curl that braces the elbow on the inner thigh and squeezes hard at the top for a strong peak contraction.",
            'Hammer Curl': "hammer curl with a neutral, thumbs-up grip that drives the brachialis and brachioradialis along with the biceps.",
            'Reverse Curl': "reverse curl with an overhand, palms-down grip that shifts the work onto the brachioradialis and wrist extensors.",
            'Zottman Curl': "Zottman curl that lifts palms-up, rotates to palms-down at the top, then lowers under control to train the biceps up and the forearm down.",
        }.get(mv, "biceps curl that lifts the load under control with the upper arms still to isolate the elbow flexors.")
    else:
        rl = {
            'Pushdown': "cable pushdown that hammers the lateral head of the triceps, extending the weight down with the elbows pinned to the sides.",
            'Overhead Extension': "overhead extension that biases the long head of the triceps through a deep stretch behind the head.",
            'Lying Extension': "lying extension, or skullcrusher, that loads the triceps under stretch by lowering toward the forehead and then extending.",
            'Kickback': "kickback that peaks tension on the triceps at full lockout with the upper arm fixed parallel to the floor.",
            'Close-Grip Press': "narrow, close-grip press that builds triceps mass while the chest and front delts assist.",
            'Dip': "bodyweight dip that drives the triceps hard, with the chest and front delts helping you lower and lock out.",
        }.get(mv, "isolation movement that extends the elbow to work the triceps.")
    if mv == 'Stretch/Mobility':
        rl = "arm stretch and mobility drill that lengthens the biceps and triceps and opens the shoulder."
    # rl begins with a bare noun-phrase; add exactly ONE article in front of it.
    art = 'An ' if rl[0] in 'AEIOUaeiou' else 'A '
    eqs = '' if eq == 'Bodyweight' else f" Performed with {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}."
    return art + rl + eqs


def instructions(eq, mv):
    """Step-by-step how-to, written to be READ ALOUD (speaker/TTS). Returns a list.
    Triceps text is reused from the triceps builder; biceps text is written here in
    the same setup -> execution -> squeeze -> return -> form-cue style."""
    bar = ('the bar' if 'Bar' in eq or eq in ('Barbell', 'Smith Machine', 'EZ / SZ Bar') else
           'the rope' if eq == 'Cable' else 'the weight' if eq in ('Dumbbell', 'Kettlebell', 'Household / Improvised') else 'the handle')
    # biceps "tool" wording: cable -> handle, free weights -> weight, bars -> bar
    btool = ('the bar' if 'Bar' in eq or eq in ('Barbell', 'Smith Machine', 'EZ / SZ Bar') else
             'the handle' if eq == 'Cable' else 'the weight' if eq in ('Dumbbell', 'Kettlebell', 'Household / Improvised') else 'the handles')

    BICEPS_STEPS = {
        'Curl': [
            f"Stand tall holding {btool} with your palms facing up and your arms hanging straight down.",
            "Pin your elbows to your sides and brace your core — your upper arms stay still the whole set.",
            f"Curl {btool} up toward your shoulders from a full stretch to a hard squeeze, contracting your biceps at the top.",
            "Hold the squeeze for a moment at the top.",
            f"Lower {btool} slowly under control until your arms are fully straight. Keep your elbows pinned — don't swing or use your back.",
        ],
        'Incline Curl': [
            f"Sit back on an incline bench so your arms hang straight down behind your torso, holding {btool} with your palms up.",
            "Let your shoulders relax back and keep your upper arms still — this stretch is what biases the long head.",
            f"Curl {btool} up toward your shoulders without letting your elbows drift forward.",
            "Squeeze your biceps hard at the top for a count.",
            f"Lower {btool} slowly back to the full stretch, keeping the tension on your biceps the whole way.",
        ],
        'Drag Curl': [
            f"Stand tall holding {btool} with your palms up and your arms hanging straight.",
            "Brace your core and get ready to pull your elbows back rather than forward.",
            f"Curl {btool} up while dragging it close along your body, sliding your elbows behind you.",
            "Squeeze your biceps at the top, keeping the bar near your torso.",
            f"Lower {btool} slowly back down the same path under control. Keep it close — don't let it swing away from you.",
        ],
        'Preacher Curl': [
            f"Set your upper arms flat on the preacher pad and hold {btool} with your palms up and arms extended.",
            "Keep your arms fixed on the pad — the angle kills momentum and loads the short head.",
            f"Curl {btool} up from the deep stretch toward your shoulders, squeezing your biceps.",
            "Hold the contraction briefly at the top.",
            f"Lower {btool} slowly until your arms are almost straight, staying in control. Never bounce out of the bottom stretch.",
        ],
        'Spider Curl': [
            f"Lie chest-down on an incline bench so your arms hang straight down, holding {btool} with your palms up.",
            "Let your arms hang dead straight and keep your upper arms vertical and still.",
            f"Curl {btool} straight up toward your shoulders using only your biceps, with no body english.",
            "Squeeze hard at the top for a moment.",
            f"Lower {btool} slowly back to the hanging stretch, keeping constant tension on the biceps.",
        ],
        'Concentration Curl': [
            "Sit on a bench, lean forward, and brace the back of your working arm against the inside of your thigh.",
            f"Let {btool} hang straight down with your palm facing away and your elbow fixed against your thigh.",
            f"Curl {btool} up toward your shoulder, keeping your upper arm locked against your thigh.",
            "Squeeze the peak of your biceps hard at the top.",
            f"Lower {btool} slowly back to a full stretch under control, then repeat before switching arms.",
        ],
        'Hammer Curl': [
            f"Stand tall holding {btool} with a neutral, thumbs-up grip and your arms hanging straight.",
            "Pin your elbows to your sides and keep your wrists straight in that neutral position.",
            f"Curl {btool} up toward your shoulders, keeping your thumbs pointing up the whole way.",
            "Squeeze your biceps and forearms at the top for a count.",
            f"Lower {btool} slowly back down under control, keeping your upper arms still.",
        ],
        'Reverse Curl': [
            f"Stand tall holding {btool} with an overhand, palms-down grip and your arms hanging straight.",
            "Pin your elbows to your sides and keep your wrists firm and straight.",
            f"Curl {btool} up toward your shoulders, leading with the backs of your hands so the work lands on your forearms.",
            "Squeeze the brachioradialis and biceps at the top.",
            f"Lower {btool} slowly back down under control, keeping your elbows pinned and your wrists locked.",
        ],
        'Zottman Curl': [
            f"Stand tall holding {btool} with your palms up and your arms hanging straight at your sides.",
            f"Curl {btool} up toward your shoulders with your palms up, pinning your elbows to your sides.",
            "At the top, rotate your wrists so your palms now face down.",
            f"Lower {btool} slowly with the palms-down grip to load the forearms on the way down.",
            "Rotate back to palms-up at the bottom and repeat. Control the lowering phase — that's where the work is.",
        ],
    }
    TRICEPS_STEPS = {
        'Pushdown': [
            f"Stand tall facing the cable and grip {bar} with your hands at about chest height, elbows tucked in close to your sides.",
            "Pin your upper arms against your body and brace your core — these stay still the whole set.",
            f"Extend your elbows and push {bar} straight down until your arms are fully locked out.",
            "Squeeze your triceps hard at the bottom for a moment.",
            f"Let {bar} rise back up slowly under control, stopping when your forearms reach about parallel. Keep your elbows pinned — only your forearms should move.",
        ],
        'Overhead Extension': [
            f"Set up tall with {bar} raised overhead and your arms fully extended, elbows pointing forward.",
            "Keep your elbows high and close to your head — this is the key to working the long head.",
            f"Bend at the elbows and lower {bar} behind your head until you feel a deep stretch in your triceps.",
            f"Drive {bar} back up to full lockout by straightening your arms, keeping your upper arms still.",
            "Keep your ribs down and core braced so your lower back doesn't arch.",
        ],
        'Lying Extension': [
            f"Lie back on the bench and press {bar} up so your arms are straight and your upper arms are vertical.",
            "Tip your arms back slightly so the load stays over your triceps, not your shoulders.",
            f"Bend only at the elbows and lower {bar} toward your forehead or just behind your head, keeping your upper arms still.",
            f"Extend your elbows to press {bar} back up to the start, squeezing your triceps at the top.",
            "Move slowly and keep your upper arms locked in place the entire time — don't let your elbows flare.",
        ],
        'Close-Grip Press': [
            f"Lie back and take a narrow grip on {bar}, hands roughly shoulder-width apart.",
            f"Unrack {bar} and hold it over your chest with your arms straight.",
            f"Lower {bar} toward your lower chest, keeping your elbows tucked in close to your body.",
            f"Press {bar} back up by extending your elbows, driving with your triceps until your arms lock out.",
            "Keep your wrists stacked over your elbows and your elbows tucked — don't let them flare wide.",
        ],
        'Dip': [
            "Grip the bars and press up to support yourself with your arms straight and your body fairly upright.",
            "Stay tall and keep your torso vertical to bias the triceps, with only a slight forward lean.",
            "Bend your elbows and lower your body until they reach about ninety degrees, keeping your elbows close to your sides.",
            "Press back up by straightening your arms until you lock out at the top.",
            "Keep your shoulders down away from your ears and control the descent — don't drop fast.",
        ],
        'Kickback': [
            f"Hinge forward at the hips with a flat back and hold {bar}, then raise your upper arm until it's parallel to the floor.",
            "Pin that upper arm against your side — it stays fixed for every rep.",
            f"Extend your elbow and drive {bar} back until your arm is fully straight.",
            "Squeeze your triceps hard at full lockout for a count.",
            f"Lower {bar} slowly back to a ninety-degree bend, keeping your upper arm still and parallel the whole time.",
        ],
    }
    STRETCH = [
        "Move into the stretch slowly until you feel a gentle pull through your upper arm.",
        "Hold the position and breathe deeply, letting the muscle lengthen with each exhale.",
        "Ease out of the stretch under control and repeat on the other arm if needed.",
    ]
    if mv == 'Stretch/Mobility':
        return STRETCH
    if is_biceps(mv):
        return BICEPS_STEPS.get(mv, [
            f"Stand tall holding {btool} with your palms up and your elbows pinned to your sides.",
            f"Curl {btool} up to a hard squeeze, keeping your upper arms still.",
            f"Lower {btool} slowly under control to a full stretch and repeat.",
        ])
    return TRICEPS_STEPS.get(mv, [
        f"Set up with your elbows fixed and a firm grip on {bar}.",
        f"Extend your elbows to move {bar}, squeezing your triceps at lockout.",
        f"Return {bar} slowly under control and repeat, keeping your upper arms still.",
    ])


def folder_for(gender, mv):
    if mv == 'Stretch/Mobility':
        cat = 'Stretch & Mobility'
    elif is_biceps(mv):
        reg = biceps_region(mv)
        if mv in ('Preacher Curl', 'Concentration Curl'):
            cat = 'Biceps (Preacher & Concentration)'
        elif reg == 'Brachialis':                      # Hammer / Reverse / Zottman
            cat = 'Biceps (Hammer & Reverse)'
        else:                                          # Curl / Incline / Drag / Spider
            cat = 'Biceps (Curls)'
    else:
        cat = {
            'Pushdown': 'Triceps (Pushdowns)',
            'Overhead Extension': 'Triceps (Overhead & Lying Extensions)',
            'Lying Extension': 'Triceps (Overhead & Lying Extensions)',
            'Close-Grip Press': 'Triceps (Press & Dips)',
            'Dip': 'Triceps (Press & Dips)',
            'Kickback': 'Triceps (Kickbacks)',
        }.get(mv, 'Triceps (Pushdowns)')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 in', SRC, '- already organized.'); return
    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['side'] = 'Biceps' if is_biceps(r['movement']) else ('Stretch' if r['movement'] == 'Stretch/Mobility' else 'Triceps')
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'])
        r['flag'] = ''
        r['folder'] = folder_for(r['gender'], r['movement'])
        recs.append(r)

    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    # ---- styled xlsx ----
    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Group',
            'Primary Muscle', 'Secondary Muscles', 'Description',
            'Step-by-step Instructions (TTS)', 'Variants', 'ID', 'Video File(s)', 'Folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r['name'], r['equipment'], r['gender'], r['level'], r['movement'], r['side'],
                   r['primary'], r['secondary'], r['description'],
                   '\n'.join(f"{i+1}. {s}" for i, s in enumerate(r['instructions'])),
                   r['variants'], r['id'], r['files'], r['folder']])
    ARIAL = 'Arial'
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'), 'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([40, 18, 9, 13, 20, 10, 34, 38, 58, 70, 9, 11, 50, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 7, 8, 9, 10, 13, 14))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises (deduped)', len(catalog)]); sm.append([])
    for t, k in [('By group (Biceps vs Triceps)', 'side'), ('By equipment', 'equipment'),
                 ('By gender', 'gender'), ('By level', 'level'), ('By movement type', 'movement')]:
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    sm.column_dimensions['A'].width = 40; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    # ---- organize files into Gender/Category ----
    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\nSRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\nprint('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('GROUP', dict(Counter(r['side'] for r in catalog)))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
