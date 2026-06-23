"""
Calves catalog builder — same template as back/chest/biceps but with CALVES
exercise science (gastrocnemius / soleus / tibialis), and a step-by-step
`instructions` array per exercise written to be READ ALOUD by the app's
speaker/TTS feature. Run from this folder: python build_calves_catalog.py
"""
import os, re, json, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Calves_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "calves_catalog.json")
TITLE = "Calves"


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)
    name_part = s.split('_')[0]
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender, 'bodypart': 'Lower Legs'}


def equipment(n):
    l = n.lower()
    if 'smith' in l: return 'Smith Machine'
    if 'barbell' in l: return 'Barbell'
    if 'dumbbell' in l: return 'Dumbbell'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'cable' in l: return 'Cable'
    if 'sled' in l or 'leg press' in l or 'leg-press' in l or 'hack' in l: return 'Leg-Press / Sled'
    if 'lever' in l or 'machine' in l or 'plate-loaded' in l: return 'Machine (Lever)'
    if re.search(r'\bband\b', l) or 'resistance band' in l: return 'Resistance Band'
    if re.search(r'\brope\b', l): return 'Rope'
    return 'Bodyweight'


def movement(n):
    l = n.lower()
    # Stretch / mobility / self-myofascial (foam-rolling). `\broll\b` catches
    # "Roll Upper Back" etc. but NOT "Rollout" (one word = no boundary = a core move).
    if any(k in l for k in ['stretch', 'mobility', 'foam']) or re.search(r'\broll\b', l):
        return 'Stretch/Mobility'
    if 'seated' in l: return 'Seated Raise'
    if 'donkey' in l: return 'Donkey Raise'
    if 'leg press' in l or 'leg-press' in l or 'press' in l: return 'Leg-Press Raise'
    if 'standing' in l or 'raise' in l or 'calf' in l: return 'Standing Raise'
    return 'Standing Raise'


CATEGORY = {  # movement -> folder category (NO "/" — use "&")
    'Standing Raise': 'Standing (Gastrocnemius)',
    'Donkey Raise': 'Standing (Gastrocnemius)',
    'Seated Raise': 'Seated (Soleus)',
    'Leg-Press Raise': 'Leg-Press & Machine',
    'Stretch/Mobility': 'Stretch & Mobility',
}

ADV = ['one-arm', 'one arm', 'single-arm', 'single arm', 'one-leg', 'one leg', 'single-leg', 'single leg', 'deficit']
BEG = ['machine', 'lever', 'seated', 'leg press', 'leg-press', 'sled', 'hack', 'stretch', 'mobility']


def level(n, mv):
    l = n.lower()
    if mv == 'Stretch/Mobility': return 'Beginner'
    if any(k in l for k in ['single-leg', 'single leg', 'one-leg', 'one leg', 'deficit']): return 'Advanced'
    if any(k in l for k in BEG): return 'Beginner'
    return 'Intermediate'


def muscles(mv):
    return {
        'Standing Raise':   ('Gastrocnemius', 'Soleus, Tibialis posterior'),
        'Donkey Raise':     ('Gastrocnemius', 'Soleus, Tibialis posterior'),
        'Seated Raise':     ('Soleus', 'Gastrocnemius, Tibialis posterior'),
        'Leg-Press Raise':  ('Gastrocnemius & Soleus', 'Tibialis posterior'),
        'Stretch/Mobility': ('Gastrocnemius & Soleus (stretch)', 'Achilles tendon, Plantar fascia'),
    }.get(mv, ('Gastrocnemius', 'Soleus, Tibialis posterior'))


def describe(eq, mv):
    rl = {
        'Standing Raise': "standing calf raise that targets the gastrocnemius — the large, visible calf muscle — through a big stretch and a high contraction with the knees straight.",
        'Donkey Raise': "bent-over donkey calf raise that stretches and loads the gastrocnemius hard across a long range of motion.",
        'Seated Raise': "seated calf raise with the knees bent at ninety degrees, which shifts the work onto the soleus — the deep calf muscle that drives ankle endurance.",
        'Leg-Press Raise': "leg-press calf raise that loads both heads of the calf — gastrocnemius and soleus — by pressing the platform with the balls of the feet only.",
        'Stretch/Mobility': "calf stretch and mobility drill that lengthens the gastrocnemius, soleus and Achilles tendon to improve ankle range of motion.",
    }.get(mv, "calf-raising movement for the lower legs.")
    art = 'An ' if rl[0] in 'AEIOUaeiou' else 'A '
    eqs = '' if eq == 'Bodyweight' else f" Performed with {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}."
    return art + rl + eqs


def instructions(eq, mv):
    """Step-by-step how-to, written to be READ ALOUD (speaker/TTS). Returns a list."""
    tool = 'the bar' if 'Bar' in eq or eq in ('Barbell', 'Smith Machine') else \
           'the handle' if eq == 'Cable' else 'the weight' if eq in ('Dumbbell', 'Kettlebell') else 'the load'
    STEPS = {
        'Standing Raise': [
            f"Stand tall on the balls of your feet on a step or block, with {tool} loading your body and your heels hanging free.",
            "Keep your knees straight and your core braced, then rise as high as possible onto your toes.",
            "Squeeze your calves hard at the very top and pause there for a full second.",
            "Lower under control all the way down to a deep stretch with your heels well below parallel.",
            "Pause for a moment in that bottom stretch, then drive back up. Never bounce out of the bottom.",
        ],
        'Donkey Raise': [
            f"Bend forward at the hips until your torso is roughly parallel to the floor, with {tool} resting across your hips and the balls of your feet on a block.",
            "Keep your knees straight and your back flat, then rise as high as possible onto your toes.",
            "Squeeze your calves at the top and hold the pause for a full second.",
            "Lower slowly into a deep stretch with your heels dropping well below the block.",
            "Pause at the bottom stretch, then press back up under control through the full range.",
        ],
        'Seated Raise': [
            f"Sit with your knees bent at ninety degrees and the pad resting low across your thighs, with the balls of your feet on the platform.",
            "Drive through the balls of your feet and rise up onto your toes as high as you can.",
            "Squeeze the soleus hard at the top and pause there for a full second.",
            "Lower slowly to a deep stretch with your heels dropping well below the platform.",
            "Pause in the bottom stretch, then press back up. Move slowly — this targets the soleus, so control beats speed.",
        ],
        'Leg-Press Raise': [
            "Sit in the leg press and place only the balls of your feet on the lower edge of the platform, heels hanging off.",
            "Unlock the safeties and let the platform push your toes down into a deep stretch at the bottom.",
            "Press the platform away using the balls of your feet only, extending your ankles as far as possible.",
            "Pause and squeeze hard at the top for a full second.",
            "Lower under control back to the deep bottom stretch and pause before the next rep. Keep your knees still throughout.",
        ],
        'Stretch/Mobility': [
            "Move into the calf stretch slowly until you feel a gentle pull through your calf and Achilles.",
            "Hold the position and breathe deeply, letting the calf lengthen with each exhale.",
            "Ease out of the stretch under control and repeat on the other side if needed.",
        ],
    }
    return STEPS.get(mv, [
        f"Set up on the balls of your feet with {tool} loaded and your heels free.",
        "Rise as high as possible onto your toes and pause-squeeze at the top.",
        "Lower under control to a deep stretch at the bottom and pause before the next rep.",
    ])


def folder_for(gender, mv):
    cat = CATEGORY.get(mv, 'Standing (Gastrocnemius)')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 - already organized.'); return
    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['category'] = CATEGORY.get(r['movement'], 'Standing (Gastrocnemius)')
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'])
        r['flag'] = ''
        r['folder'] = folder_for(r['gender'], r['movement'])
        recs.append(r)
    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Primary Muscle',
            'Secondary Muscles', 'Description', 'Step-by-step Instructions (TTS)', 'Variants', 'ID', 'Video File(s)', 'Folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r['name'], r['equipment'], r['gender'], r['level'], r['movement'], r['primary'],
                   r['secondary'], r['description'], '\n'.join(f"{i+1}. {s}" for i, s in enumerate(r['instructions'])),
                   r['variants'], r['id'], r['files'], r['folder']])
    ARIAL = 'Arial'
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'), 'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([38, 18, 9, 13, 20, 34, 38, 56, 70, 9, 11, 50, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 6, 7, 8, 9, 12, 13))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises', len(catalog)]); sm.append([])
    for t, k in [('By equipment', 'equipment'), ('By gender', 'gender'), ('By level', 'level'), ('By movement type', 'movement')]:
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    sm.column_dimensions['A'].width = 38; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\nSRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\nprint('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
