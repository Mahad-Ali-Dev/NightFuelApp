"""
Triceps catalog builder — same template as back/chest/biceps but with TRICEPS
exercise science (long / lateral / medial heads, anconeus), and a step-by-step
`instructions` array per exercise written to be READ ALOUD by the app's
speaker/TTS feature. Run from this folder: python build_triceps_catalog.py
"""
import os, re, json, shutil
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(SRC, "Triceps_Exercise_Catalog.xlsx")
OUT_JSON = os.path.join(SRC, "triceps_catalog.json")
TITLE = "Triceps"


def parse(fn):
    base = fn[:-4]
    m = re.match(r'^(\d+)-(.*)$', base)
    rid, s = (m.group(1), m.group(2)) if m else ('', base)
    s = re.sub(r'\s*\(\d+\)\s*$', '', s)
    name_part = s.split('_')[0]
    low = name_part.lower()
    gender = 'Female' if ('female' in low or re.search(r'\(fe', low)) else 'Male'
    name = re.sub(r'\((?:fe)?male\)', '', name_part, flags=re.I)
    name = re.sub(r'-?FIX2?', '', name, flags=re.I)
    name = re.sub(r'\(version[- ]?\d\)', '', name, flags=re.I)
    name = name.replace('-', ' ')
    name = re.sub(r'\s+', ' ', name).strip(' _-')
    return {'file': fn, 'id': rid, 'name': name, 'gender': gender, 'bodypart': 'Upper Arms'}


def equipment(n):
    l = n.lower()
    if 'smith' in l: return 'Smith Machine'
    if 'ez' in l.split() or 'ez-bar' in l or 'ez bar' in l or 'sz-bar' in l or 'sz bar' in l: return 'EZ / SZ Bar'
    if 'barbell' in l or 'olympic' in l: return 'Barbell'
    if 'dumbbell' in l: return 'Dumbbell'
    if 'kettlebell' in l: return 'Kettlebell'
    if 'cable' in l: return 'Cable'
    if 'machine' in l or 'lever' in l or 'plate-loaded' in l: return 'Machine (Lever)'
    if re.search(r'\bband\b', l) or 'resistance band' in l: return 'Resistance Band'
    if 'suspension' in l or 'suspender' in l or 'trx' in l or 'ring' in l: return 'Suspension (TRX)'
    if 'bottle' in l or 'towel' in l: return 'Household / Improvised'
    return 'Bodyweight'


def movement(n):
    l = n.lower()
    # Stretch / mobility / self-myofascial (foam-rolling). `\broll\b` catches
    # "Roll ..." etc. but NOT "Rollout" (one word = no boundary = a core move).
    if any(k in l for k in ['stretch', 'mobility', 'opener', 'foam', 'roller',
                            'release', 'decompress', 'dead hang']) or re.search(r'\broll\b', l):
        return 'Stretch/Mobility'
    if 'kickback' in l: return 'Kickback'
    if 'skull' in l or 'lying' in l or 'french' in l: return 'Lying Extension'
    if 'overhead' in l or 'seated extension' in l: return 'Overhead Extension'
    if 'pushdown' in l or 'press-down' in l or 'pressdown' in l: return 'Pushdown'
    if 'dip' in l: return 'Dip'
    if ('close' in l and 'grip' in l) or 'jm ' in l or 'diamond' in l or 'narrow' in l: return 'Close-Grip Press'
    if 'extension' in l: return 'Overhead Extension'
    return 'Pushdown'


REGION = {  # movement -> training category (folder, NO "/" — use "&")
    'Pushdown': 'Pushdowns',
    'Overhead Extension': 'Overhead Extensions',
    'Lying Extension': 'Lying Extensions (Skullcrusher)',
    'Close-Grip Press': 'Press & Dips',
    'Dip': 'Press & Dips',
    'Kickback': 'Kickbacks',
    'Stretch/Mobility': 'Stretch & Mobility',
}

ADV = ['one-arm', 'one arm', 'single-arm', 'single arm', 'deficit', 'weighted dip',
       'handstand', 'ring', 'side-lying', 'side lying', 'single-leg', 'single leg']
BEG = ['assisted', 'machine', 'lever', 'cable', 'band', 'bench dip', 'wall']


def level(n, mv):
    l = n.lower()
    if mv == 'Stretch/Mobility': return 'Beginner'
    if any(k in l for k in ADV): return 'Advanced'
    if any(k in l for k in BEG): return 'Beginner'
    return 'Intermediate'


def muscles(mv):
    return {
        'Overhead Extension':  ('Triceps brachii (long head)', 'Anconeus'),
        'Pushdown':            ('Triceps brachii (lateral head)', 'Anconeus'),
        'Lying Extension':     ('Triceps brachii (long & lateral head)', 'Anconeus'),
        'Kickback':            ('Triceps brachii (lateral & long head)', 'Anconeus'),
        'Close-Grip Press':    ('Triceps brachii', 'Pectoralis major, Anterior deltoid'),
        'Dip':                 ('Triceps brachii', 'Pectoralis major (lower), Anterior deltoid'),
        'Stretch/Mobility':    ('Triceps brachii (stretch)', 'Latissimus dorsi, Posterior deltoid'),
    }.get(mv, ('Triceps brachii', 'Anconeus'))


def describe(eq, mv):
    rl = {
        'Pushdown': "a cable isolation that hammers the lateral head of the triceps — elbows pinned to your sides as you extend the weight down.",
        'Overhead Extension': "an overhead isolation that biases the LONG head of the triceps, working it through a deep stretch behind the head.",
        'Lying Extension': "a lying extension (skullcrusher) that loads the triceps under stretch — lowering toward the forehead, then extending.",
        'Kickback': "an isolation lift that peaks tension on the triceps at full lockout, with the upper arm fixed parallel to the floor.",
        'Close-Grip Press': "a narrow-grip press that builds triceps mass while the chest and front delts assist.",
        'Dip': "a bodyweight press that drives the triceps hard, with the chest and front delts helping you lower and lock out.",
        'Stretch/Mobility': "a triceps stretch/mobility drill to lengthen the triceps and open the shoulder overhead.",
    }.get(mv, "an isolation movement for the triceps.")
    # rl already begins with its own article ("a ..." / "an ..."), so just capitalize it.
    eqs = '' if eq == 'Bodyweight' else f" Performed with {('an ' if eq[0] in 'AEIOU' else 'a ')}{eq.lower()}."
    return rl[0].upper() + rl[1:] + eqs


def instructions(eq, mv):
    """Step-by-step how-to, written to be READ ALOUD (speaker/TTS). Returns a list."""
    tool = 'the bar' if 'Bar' in eq or eq in ('Barbell', 'Smith Machine', 'EZ / SZ Bar') else \
           'the rope' if eq == 'Cable' else 'the weight' if eq in ('Dumbbell', 'Kettlebell', 'Household / Improvised') else 'the handle'
    STEPS = {
        'Pushdown': [
            f"Stand tall facing the cable and grip {tool} with your hands at about chest height, elbows tucked in close to your sides.",
            "Pin your upper arms against your body and brace your core — these stay still the whole set.",
            f"Extend your elbows and push {tool} straight down until your arms are fully locked out.",
            "Squeeze your triceps hard at the bottom for a moment.",
            f"Let {tool} rise back up slowly under control, stopping when your forearms reach about parallel. Keep your elbows pinned — only your forearms should move.",
        ],
        'Overhead Extension': [
            f"Set up tall with {tool} raised overhead and your arms fully extended, elbows pointing forward.",
            "Keep your elbows high and close to your head — this is the key to working the long head.",
            f"Bend at the elbows and lower {tool} behind your head until you feel a deep stretch in your triceps.",
            f"Drive {tool} back up to full lockout by straightening your arms, keeping your upper arms still.",
            "Keep your ribs down and core braced so your lower back doesn't arch.",
        ],
        'Lying Extension': [
            f"Lie back on the bench and press {tool} up so your arms are straight and your upper arms are vertical.",
            "Tip your arms back slightly so the load stays over your triceps, not your shoulders.",
            f"Bend only at the elbows and lower {tool} toward your forehead or just behind your head, keeping your upper arms still.",
            f"Extend your elbows to press {tool} back up to the start, squeezing your triceps at the top.",
            "Move slowly and keep your upper arms locked in place the entire time — don't let your elbows flare.",
        ],
        'Close-Grip Press': [
            f"Lie back and take a narrow grip on {tool}, hands roughly shoulder-width apart.",
            f"Unrack {tool} and hold it over your chest with your arms straight.",
            f"Lower {tool} toward your lower chest, keeping your elbows tucked in close to your body.",
            f"Press {tool} back up by extending your elbows, driving with your triceps until your arms lock out.",
            "Keep your wrists stacked over your elbows and your elbows tucked — don't let them flare wide.",
        ],
        'Dip': [
            "Grip the bars and press up to support yourself with your arms straight and your body fairly upright.",
            "Stay tall and keep your torso vertical to bias the triceps, with only a slight forward lean.",
            "Bend your elbows and lower your body until they reach about ninety degrees, keeping your elbows close to your sides.",
            "Press back up by straightening your arms until you lock out at the top.",
            "Keep your shoulders down away from your ears and control the descent — don't drop fast.",
        ],
        'Kickback': [
            f"Hinge forward at the hips with a flat back and hold {tool}, then raise your upper arm until it's parallel to the floor.",
            "Pin that upper arm against your side — it stays fixed for every rep.",
            f"Extend your elbow and drive {tool} back until your arm is fully straight.",
            "Squeeze your triceps hard at full lockout for a count.",
            f"Lower {tool} slowly back to a ninety-degree bend, keeping your upper arm still and parallel the whole time.",
        ],
        'Stretch/Mobility': [
            "Move into the stretch slowly until you feel a gentle pull along the back of your upper arm.",
            "Hold the position and breathe deeply, letting the triceps lengthen with each exhale.",
            "Ease out of the stretch under control and repeat on the other arm if needed.",
        ],
    }
    return STEPS.get(mv, [
        f"Set up with your elbows fixed and a firm grip on {tool}.",
        f"Extend your elbows to move {tool}, squeezing your triceps at lockout.",
        f"Return {tool} slowly under control and repeat, keeping your upper arms still.",
    ])


def folder_for(gender, mv):
    cat = REGION.get(mv, 'Pushdowns')
    return gender + '/' + cat


def main():
    files = sorted(f for f in os.listdir(SRC) if f.lower().endswith('.mp4'))
    if not files:
        print('No top-level .mp4 - already organized.'); return
    recs = []
    for f in files:
        r = parse(f)
        r['equipment'] = equipment(r['name'])
        r['movement'] = movement(r['name'])
        r['region'] = REGION.get(r['movement'], 'Pushdowns')
        r['level'] = level(r['name'], r['movement'])
        r['primary'], r['secondary'] = muscles(r['movement'])
        r['description'] = describe(r['equipment'], r['movement'])
        r['instructions'] = instructions(r['equipment'], r['movement'])
        r['flag'] = ''
        r['folder'] = folder_for(r['gender'], r['movement'])
        recs.append(r)
    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'].lower(), r['gender'])].append(r)
    catalog = []
    for grp in groups.values():
        b = dict(grp[0]); b['variants'] = len(grp)
        b['files'] = '; '.join(sorted(set(g['file'] for g in grp)))
        catalog.append(b)
    catalog.sort(key=lambda r: (r['gender'], r['folder'], r['equipment'], r['name']))
    json.dump(catalog, open(OUT_JSON, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)

    wb = Workbook(); ws = wb.active; ws.title = f'{TITLE} Catalog'
    cols = ['Exercise Name', 'Equipment', 'Gender', 'Level', 'Movement Type', 'Primary Muscle',
            'Secondary Muscles', 'Description', 'Step-by-step Instructions (TTS)', 'Variants', 'ID', 'Video File(s)', 'Folder']
    ws.append(cols)
    for r in catalog:
        ws.append([r['name'], r['equipment'], r['gender'], r['level'], r['movement'], r['primary'],
                   r['secondary'], r['description'], '\n'.join(f"{i+1}. {s}" for i, s in enumerate(r['instructions'])),
                   r['variants'], r['id'], r['files'], r['folder']])
    ARIAL = 'Arial'
    thin = Side(style='thin', color='D9D9D9'); border = Border(thin, thin, thin, thin)
    lvl_fill = {'Beginner': PatternFill('solid', fgColor='C6EFCE'), 'Intermediate': PatternFill('solid', fgColor='FFEB9C'),
                'Advanced': PatternFill('solid', fgColor='FFC7CE')}
    for i, w in enumerate([38, 18, 9, 13, 20, 34, 38, 56, 70, 9, 11, 50, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row, c); cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 6, 7, 8, 9, 12, 13))); cell.border = border
        lv = ws.cell(row, 4); lv.fill = lvl_fill.get(lv.value, PatternFill())
        lv.alignment = Alignment(horizontal='center', vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    sm = wb.create_sheet('Summary')
    sm.append([f'{TITLE} Exercise Catalog - Summary']); sm['A1'].font = Font(name=ARIAL, bold=True, size=14)
    sm.append([]); sm.append(['Total video files', len(files)]); sm.append(['Unique exercises', len(catalog)]); sm.append([])
    for t, k in [('By equipment', 'equipment'), ('By gender', 'gender'), ('By level', 'level'), ('By movement type', 'movement')]:
        sm.append([t]); sm.cell(sm.max_row, 1).font = Font(name=ARIAL, bold=True, size=11)
        for val, cnt in sorted(Counter(r[k] for r in catalog).items(), key=lambda x: -x[1]):
            sm.append([val, cnt])
        sm.append([])
    sm.column_dimensions['A'].width = 38; sm.column_dimensions['B'].width = 12
    wb.save(OUT_XLSX)

    manifest = []
    for f in files:
        r = parse(f)
        rel = folder_for(r['gender'], movement(r['name']))
        dest_dir = os.path.join(SRC, *rel.split('/')); os.makedirs(dest_dir, exist_ok=True)
        src = os.path.join(SRC, f); dst = os.path.join(dest_dir, f)
        if os.path.exists(src) and os.path.abspath(src) != os.path.abspath(dst):
            shutil.move(src, dst)
        manifest.append({'file': f, 'folder': rel})
    json.dump(manifest, open(os.path.join(SRC, '_organize_manifest.json'), 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
    open(os.path.join(SRC, '_undo_organize.py'), 'w', encoding='utf-8').write(
        "import os, json, shutil\nSRC=os.path.dirname(os.path.abspath(__file__))\n"
        "for e in json.load(open(os.path.join(SRC,'_organize_manifest.json'),encoding='utf-8')):\n"
        "    cur=os.path.join(SRC,*e['folder'].split('/'),e['file'])\n"
        "    if os.path.exists(cur): shutil.move(cur, os.path.join(SRC,e['file']))\nprint('undone')\n")
    print('FILES', len(files), 'UNIQUE', len(catalog))
    print('LEVELS', dict(Counter(r['level'] for r in catalog)))
    print('CATEGORIES', dict(Counter(r['folder'] for r in catalog)))


if __name__ == '__main__':
    main()
